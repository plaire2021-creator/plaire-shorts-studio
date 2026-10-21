# -*- coding: utf-8 -*-
"""
통합파일(xlsx) → 구글 시트 '그대로' 반영 + 전체 칸 대조
------------------------------------------------------------------
구글 시트(주소·공유 그대로)의 내용을 통합 엑셀 파일로 통째로 교체한다(Drive 업로드 변환).
탭 이름·병합·색·열 너비·차트까지 엑셀과 같아진다. 이후 모든 탭의 모든 칸을 엑셀과 대조한다.

  ⚠ 구글 시트에서 직접 고친 내용은 교체 때 사라진다 → 수정은 통합 엑셀에서 할 것.

실행:
  python gsheet_mirror.py --company 가즈다가구 --file "통합/2026_가즈다가구 주간판매데이터 (10월12주차).xlsx"
  python gsheet_mirror.py --company 가즈다가구 --file ... --verify-only
"""
import argparse
import os
import sys
from pathlib import Path

import openpyxl

import master_book as mb

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

SHEET_IDS = {
    "플레르가구": "1caviDCetj3RsdhRmgqRarfK7y6ee8Kt7t9mDe08VmU4",
    "가즈다가구": "1VoCNP7TJ01JTa0uR8XDM1jcuumyU4ts6A0yrXfJhXgo",
}
KEY = Path.home() / ".secrets" / "google-sheets-key.json"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload_replace(company, path):
    from google.oauth2 import service_account
    from google.auth.transport.requests import AuthorizedSession
    creds = service_account.Credentials.from_service_account_file(
        str(KEY), scopes=["https://www.googleapis.com/auth/drive"])
    s = AuthorizedSession(creds)
    fid = SHEET_IDS[company]
    with open(path, "rb") as fh:
        r = s.patch(f"https://www.googleapis.com/upload/drive/v3/files/{fid}?uploadType=media&supportsAllDrives=true"
                    "&fields=id,name,mimeType,modifiedTime",
                    data=fh.read(), headers={"Content-Type": XLSX}, timeout=300)
    if r.status_code != 200:
        raise SystemExit(f"구글 시트 교체 실패({r.status_code}): {r.text[:500]}")
    meta = r.json()
    if meta.get("mimeType") != "application/vnd.google-apps.spreadsheet":
        raise SystemExit(f"구글 시트 형식이 아님: {meta}")
    return meta


BROKEN = "#원본오류수식"


def _excel_value(ws, v):
    if isinstance(v, str) and v.startswith("="):
        e = mb.eval_cell(ws, v)
        if e is None:
            return BROKEN                   # 원본에서도 계산 안 되는 수식(예: =L11:P2186635326093 → #NAME?)
        v = e
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, str):
        v = v.strip()
        if v == "":
            return None
    return v


def _sheet_value(v):
    if v in ("", None):
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip() or None
    return v


def verify_mirror(company, path, rep=None):
    """엑셀 전 탭·전 칸 ↔ 구글 시트(서식 없는 원래 값) 대조."""
    from sheets_client import open_sheet
    import verify_master as vm
    rep = rep or vm.Rep()
    wb = openpyxl.load_workbook(path)
    sh = open_sheet("https://docs.google.com/spreadsheets/d/" + SHEET_IDS[company])
    gtabs = {w.title: w for w in sh.worksheets()}
    meta = sh.fetch_sheet_metadata({"fields": "sheets(properties(title,hidden),merges)"})
    gmerge = {x["properties"]["title"]: len(x.get("merges", [])) for x in meta["sheets"]}
    ghidden = {x["properties"]["title"]: bool(x["properties"].get("hidden")) for x in meta["sheets"]}
    rep.sec("구글 시트 ↔ 통합 엑셀(전 탭·전 칸)")
    broken = []
    rep.ok(list(gtabs) == wb.sheetnames, f"탭 구성: 시트 {list(gtabs)} / 엑셀 {wb.sheetnames}")
    for ws in wb.worksheets:
        g = gtabs.get(ws.title)
        if g is None:
            rep.ok(False, f"시트에 탭 없음: {ws.title}")
            continue
        rep.ok(ghidden.get(ws.title) == (ws.sheet_state != "visible"), f"{ws.title} 숨김 상태 다름")
        gv = g.get_values(value_render_option="UNFORMATTED_VALUE")
        n = 0
        bad = []
        for r in range(1, max(ws.max_row, len(gv)) + 1):
            grow = gv[r - 1] if r - 1 < len(gv) else []
            for c in range(1, max(ws.max_column, len(grow)) + 1):
                a = _excel_value(ws, ws.cell(r, c).value) if r <= ws.max_row and c <= ws.max_column else None
                b = _sheet_value(grow[c - 1] if c - 1 < len(grow) else None)
                if a is None and b is None:
                    continue
                n += 1
                same = a == b or (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 0.5)
                if a == BROKEN:
                    same = isinstance(b, str) and b.startswith("#")
                    broken.append(f"{ws.title}!{ws.cell(r, c).coordinate}")
                if not same:
                    bad.append(f"{ws.cell(r, c).coordinate}: 엑셀 {a!r} / 시트 {b!r}")
        rep.secs[-1][1] += n - 1 if n else 0          # 칸 수만큼 대조 건수 반영
        rep.ok(not bad, f"{ws.title}: {len(bad)}칸 다름 {bad[:5]}")
        if bad:
            rep.secs[-1][2][-1] = f"{ws.title}: {len(bad)}칸 다름 " + "; ".join(bad[:8])
        mg_x = len(ws.merged_cells.ranges)
        rep.ok(mg_x == gmerge.get(ws.title), f"{ws.title} 병합 수: 엑셀 {mg_x} / 시트 {gmerge.get(ws.title)}")
    if broken:
        print(f"  (참고) 원본부터 오류인 수식 {len(broken)}칸: {broken} — 시트에서도 오류로 표시되는지만 확인")
    return rep


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--company", required=True, choices=list(SHEET_IDS))
    ap.add_argument("--file", required=True)
    ap.add_argument("--verify-only", action="store_true")
    a = ap.parse_args()
    if not a.verify_only:
        meta = upload_replace(a.company, a.file)
        print(f"[교체 완료] {meta.get('name')} ({meta.get('modifiedTime')})")
    sys.exit(1 if verify_mirror(a.company, a.file).show() else 0)


if __name__ == "__main__":
    main()
