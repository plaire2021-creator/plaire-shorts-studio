# -*- coding: utf-8 -*-
"""
주차별 누적 스프레드시트 빌더
------------------------------------------------------------------
프로젝트 폴더의 주간판매데이터 파일들(2026_{회사} 주간판매데이터 (N월 M주차).xlsx)을
자동으로 찾아, 회사별 누적 파일(누적_{회사} 주간매출.xlsx)을 만든다.
  - 매출데이터 탭: 채널(행) × 주차(열). 왼쪽=오래된 주차, 오른쪽으로 누적.
  - 네이버/쿠팡/오늘의집 탭: 각 주차 순위표를 블록으로 왼→오 나열(A1 시작, 순위열 제거).
  - 컬러 밴드 + 테두리 그리드 + 제품 그룹 구분선.

실행:  python build_accum.py
"""
import glob
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 주간판매데이터 파일이 있는 폴더 (다른 PC면 이 경로만 바꾸면 됨)
PROJ = os.path.dirname(os.path.abspath(__file__))
COMPANIES = ["플레르가구", "가즈다가구"]

# --- 스타일 ---
GTHIN = Side(style="thin", color="808080")
GRID = Border(left=GTHIN, right=GTHIN, top=GTHIN, bottom=GTHIN)
GTOP = Side(style="medium", color="4472C4")          # 제품 그룹 구분선(윗변)
HDRF = PatternFill("solid", fgColor="D9E1F2")        # 헤더 배경
WEEKF = PatternFill("solid", fgColor="9DC3E6")       # 주차 라벨 밴드
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CEN = Alignment(horizontal="center", vertical="center")
RIT = Alignment(horizontal="right", vertical="center")
LEF = Alignment(horizontal="left", vertical="center")

# 순위 시트는 신양식 탭명(네이버/쿠팡/오늘의집)만 사용 — 구양식(…TOP20) 파일은 순위에서 제외
RANK_SHEETS = {"네이버": ["네이버"], "쿠팡": ["쿠팡"], "오늘의집": ["오늘의집"], "자사몰": ["자사몰"]}


def num(v):
    try:
        return float(str(v).replace(",", "").strip())
    except (ValueError, TypeError, AttributeError):
        return 0.0


def week_key(label):
    """'8월 24주차' → (8, 24) 정렬키."""
    m = re.search(r"(\d+)\s*월\s*(\d+)\s*주차", str(label))
    return (int(m.group(1)), int(m.group(2))) if m else (99, 99)


def weekly_files(company):
    """회사별 주간파일을 주차 오름차순으로 반환 [(라벨, 경로)…]."""
    out = []
    for f in glob.glob(os.path.join(PROJ, f"2026_{company} 주간판매데이터 (*주차).xlsx")):
        base = os.path.basename(f)
        if base.startswith("~$"):
            continue
        m = re.search(r"\(([^)]*주차)\)", base)
        if m:
            out.append((m.group(1).strip(), f))
    out.sort(key=lambda x: week_key(x[0]))
    return out


def get_sheet(wb, names):
    for n in names:
        if n in wb.sheetnames:
            return wb[n]
    return None


def build_meul(ws, company, has_cafe24):
    """매출데이터 탭: 채널(행) × 주차(열), 왼쪽=오래된 주차.
    자사몰 데이터가 없는 회사(가즈다)는 자사몰 행 생략."""
    ws["A1"] = f"{company} 주간판매 데이터"
    ws["A1"].font = Font(bold=True, size=14)
    channels = ["네이버", "쿠팡", "오늘의집"] + (["자사몰"] if has_cafe24 else [])
    sum_row = 4 + len(channels)
    for r, lab in enumerate([""] + channels + ["합계"], 3):   # 3=주차 / 채널 / 합계
        ws.cell(r, 1, lab)
    col = 2
    for label, f in weekly_files(company):
        s = openpyxl.load_workbook(f, data_only=True)["매출데이터"]
        vals = {}                                            # 라벨로 채널값 읽기(자사몰 유무 대응)
        for r in range(4, 10):
            lab = str(s.cell(r, 1).value or "").strip()
            if lab in channels:
                vals[lab] = num(s.cell(r, 2).value)
        ws.cell(3, col, label).font = Font(bold=True)
        ws.cell(3, col).fill = HDRF
        ws.cell(3, col).alignment = CEN
        tot = 0.0
        for i, ch in enumerate(channels):
            r = 4 + i
            if ch in vals:                                   # 없는 채널(자사몰 미제공 주)은 빈칸
                cell = ws.cell(r, col, int(round(vals[ch])))
                cell.number_format = "#,##0"; cell.alignment = RIT
                tot += vals[ch]
        cell = ws.cell(sum_row, col, int(round(tot)))        # 합계 행
        cell.number_format = "#,##0"; cell.alignment = RIT
        col += 1
    last = col - 1
    for cc in range(1, last + 1):
        for r in range(3, sum_row + 1):
            ws.cell(r, cc).border = BOX
        ws.column_dimensions[openpyxl.utils.get_column_letter(cc)].width = 13 if cc == 1 else 14


def build_rank(ws, company, channel):
    """순위 탭: 각 주차 순위표를 블록으로 왼→오 나열(A1 시작, 순위열 제거, 컬러+테두리)."""
    col0 = 1
    for label, f in weekly_files(company):
        wb = openpyxl.load_workbook(f, data_only=True)
        src = get_sheet(wb, RANK_SHEETS[channel])
        if src is None:
            continue                               # 구양식 파일 등은 건너뜀
        ncol = 0
        for c in range(1, src.max_column + 1):
            if src.cell(3, c).value not in (None, ""):
                ncol = c
        if ncol < 2:
            continue
        bw = ncol - 1                              # 순위열(1) 제외 블록 너비
        last_row = src.max_row - 1                 # 1행 위로 당김
        group_starts = set()
        for r in range(4, src.max_row + 1):
            if src.cell(r, 2).value not in (None, ""):
                group_starts.add(r - 1)
            for c in range(2, ncol + 1):
                v = src.cell(r, c).value
                if v is not None:
                    ws.cell(r - 1, col0 + (c - 2), v)
        ws.cell(1, col0, label)
        for br in range(1, last_row + 1):
            for bc in range(col0, col0 + bw):
                cell = ws.cell(br, bc)
                cell.border = GRID
                if br == 1:
                    cell.fill = WEEKF; cell.font = Font(bold=True); cell.alignment = CEN
                elif br == 2:
                    cell.fill = HDRF; cell.font = Font(bold=True); cell.alignment = CEN
                else:
                    if br in group_starts:
                        b = cell.border
                        cell.border = Border(left=b.left, right=b.right, top=GTOP, bottom=b.bottom)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0"; cell.alignment = RIT
                    else:
                        cell.alignment = LEF
        for c in range(2, ncol + 1):
            hdr = str(src.cell(3, c).value or "")
            w = 30 if hdr in ("제품명", "상품명") else (16 if hdr in ("옵션", "상품 코드", "옵션코드", "상품아이디") else 10)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col0 + (c - 2))].width = w
        col0 += bw + 1                             # 블록 간 1열 간격


def main():
    for company in COMPANIES:
        if not weekly_files(company):
            print(f"[건너뜀] {company}: 주간파일 없음")
            continue
        # 자사몰 데이터(순위시트)를 가진 주간파일이 하나라도 있는가 → 회사별 자사몰 표기 여부
        has_cafe24 = any(
            get_sheet(openpyxl.load_workbook(f, data_only=True), RANK_SHEETS["자사몰"])
            for _, f in weekly_files(company))
        wb = openpyxl.Workbook()
        build_meul(wb.active, company, has_cafe24)
        wb.active.title = "매출데이터"
        for ch in ("네이버", "쿠팡", "오늘의집", "자사몰"):
            if ch == "자사몰" and not has_cafe24:   # 자사몰 없는 회사는 탭 생략
                continue
            build_rank(wb.create_sheet(ch), company, ch)
        out = os.path.join(PROJ, f"누적_{company} 주간매출.xlsx")
        wb.save(out)
        wks = [lab for lab, _ in weekly_files(company)]
        print(f"[완료] {out}  (주차: {', '.join(wks)})")


if __name__ == "__main__":
    main()
