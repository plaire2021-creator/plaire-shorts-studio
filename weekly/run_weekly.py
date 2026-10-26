# -*- coding: utf-8 -*-
"""
주간 판매데이터 한 번에 처리 (플레르가구 / 가즈다가구)
------------------------------------------------------------------
원본 파일들만 넘기면 (2026-10 새 양식: 통합파일 1개에 전체 이력):
  1) 파일 종류 자동 판별 + 잘못 받은 파일 차단(네이버 기간이 한 주가 아님, 쿠팡 일자별 요약 등)
  2) 이번 주 집계(주간파일 xlsx, 중간 산출물)
  3) 원본 ↔ 주간파일 검증 (불일치면 멈춤)
  4) 통합/ 의 지난주 통합파일에 이번 주를 붙여 새 통합파일 저장 (순위 블록 오른쪽 추가, 매출데이터 B열)
  5) 통합파일 검증: 새 블록 = 주간파일, 과거 이력 1칸도 안 바뀜, 매출데이터 한 칸 이동, 차트 범위 (불일치면 멈춤)
  6) 구글 시트를 통합파일과 똑같이 교체 → 전 탭·전 칸 대조, 오류율 0% 확인

실행:
  python run_weekly.py --company 플레르가구 <원본파일들...>
  python run_weekly.py --company 플레르가구 --no-cafe24 <원본 3개>    # 그 주 자사몰 판매 없음
  python run_weekly.py --company 가즈다가구 <원본 3개>
  python run_weekly.py --check                                        # 새 PC 설치 점검
옵션: --no-sheet(시트 입력 생략) --overwrite(시트에 이미 있는 주차 다시 쓰기)
종료코드 0 = 완료(오류 0%), 그 외 = 중단/불일치.
"""
import argparse
import glob
import os
import re
import subprocess
import sys
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
COMPANIES = ("플레르가구", "가즈다가구")
HAS_CAFE24 = {"플레르가구": True, "가즈다가구": False}


def die(msg, code=2):
    print("\n[중단] " + msg)
    sys.exit(code)


def run(script, *args):
    env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    p = subprocess.run([sys.executable, os.path.join(HERE, script), *args], cwd=HERE, env=env,
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    out = "\n".join(l for l in (p.stdout + p.stderr).splitlines() if "warn" not in l.lower())
    print(out.strip())
    return p.returncode


# ---------------------------------------------------------------- 파일 판별
def classify(paths):
    import openpyxl
    import generate_weekly_sales as g
    found = {"naver": [], "coupang": [], "ohouse": [], "cafe24": []}
    for p in paths:
        if not os.path.isfile(p):
            die(f"파일이 없습니다: {p}")
        base = os.path.basename(p)
        low = base.lower()
        if low.endswith(".csv"):
            with open(p, encoding="utf-8-sig", errors="replace") as fh:
                head = fh.readline()
            if ("판매가" in head and ("주문상품명" in head or "발주일" in head)) or \
               ("판매금액" in head and ("상품번호" in head or "상품 번호" in head)):
                found["cafe24"].append(p)
                continue
            die(f"CSV 종류를 모르겠습니다(카페24 주문내역/상품별 매출 분석이 아님): {base}")
        if not low.endswith((".xlsx", ".xlsm")):
            die(f"엑셀/CSV가 아닌 파일입니다: {base}")
        if re.search(r"주간판매데이터", base):
            continue                                   # 결과 파일은 무시(base는 자동으로 찾음)
        if g.is_encrypted(p):
            found["ohouse"].append(p)
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        names = [s.lower() for s in wb.sheetnames]
        ws = wb.worksheets[0]
        hdr = [str(c.value).strip() for c in next(ws.iter_rows(min_row=1, max_row=1)) if c.value is not None]
        wb.close()
        if "daily summary metrics" in names:
            die(f"쿠팡 파일이 '일자별 요약'(DAILY_SUMMARY)입니다: {base}\n"
                "       판매 인사이트 → 상품(옵션)별 지표, 파일명 'SELLER_INSIGHTS_VENDOR_ITEM_METRICS_…' 로 다시 받아주세요.")
        if "vendor item metrics" in names or "옵션 ID" in hdr:
            found["coupang"].append(p)
        elif "sales" in names or "채널상품명" in hdr or base.startswith("sales_"):
            if "채널상품명" not in hdr:
                die(f"네이버 파일에 '채널상품명' 열이 없습니다(카테고리별로 받은 듯): {base}\n"
                    "       스마트스토어 통계 → '상품별' 판매통계로 다시 받아주세요.")
            found["naver"].append(p)
        elif "주문옵션번호" in hdr:
            found["ohouse"].append(p)
        else:
            die(f"어느 플랫폼 파일인지 모르겠습니다: {base} (시트 {wb.sheetnames})")
    for k, v in found.items():
        if len(v) > 1:
            die(f"{k} 파일이 {len(v)}개입니다. 이번 주 것 하나만 넘겨주세요: " + ", ".join(map(os.path.basename, v)))
    return {k: (v[0] if v else None) for k, v in found.items()}


def naver_week(path):
    m = re.search(r"(\d{4})(\d{2})(\d{2})-(\d{4})(\d{2})(\d{2})", os.path.basename(path))
    if not m:
        die(f"네이버 파일명에서 기간(YYYYMMDD-YYYYMMDD)을 못 읽었습니다: {os.path.basename(path)}")
    y1, m1, d1, y2, m2, d2 = map(int, m.groups())
    s, e = date(y1, m1, d1), date(y2, m2, d2)
    if (e - s).days != 6 or s.weekday() != 0:
        die(f"네이버 기간이 월~일 한 주가 아닙니다: {s} ~ {e} ({(e - s).days + 1}일)\n"
            "       스마트스토어 '상품별' 통계를 이번 주 월요일~일요일로 다시 받아주세요.")
    return y1, f"{m1}월 {d1}주차", s, e


def wkey(label):
    m = re.search(r"(\d+)\s*월\s*0*(\d+)\s*주차", label)
    return (int(m.group(1)), int(m.group(2))) if m else (0, 0)


def find_base(company, year, week):
    cands = []
    for p in glob.glob(os.path.join(HERE, f"{year}_{company} 주간판매데이터 (*주차).xlsx")):
        b = os.path.basename(p)
        m = re.search(r"\(([^)]*주차)\)", b)
        if b.startswith("~$") or not m:
            continue
        if wkey(m.group(1)) < wkey(week):
            cands.append((wkey(m.group(1)), p))
    return max(cands)[1] if cands else None


def find_master(company, year, week):
    """통합/ 폴더에서 이번 주보다 이전인 가장 최근 통합파일."""
    cands = []
    for p in glob.glob(os.path.join(HERE, "통합", f"{year}_{company} 주간판매데이터 (*주차).xlsx")):
        b = os.path.basename(p)
        m = re.search(r"\(([^)]*주차)\)", b)
        if not b.startswith("~$") and m and wkey(m.group(1)) < wkey(week):
            cands.append((wkey(m.group(1)), p))
    return max(cands)[1] if cands else None


# ---------------------------------------------------------------- 설치 점검
def check_setup():
    ok = True
    print(f"작업 폴더: {HERE}")
    for mod, pipname in [("openpyxl", "openpyxl"), ("msoffcrypto", "msoffcrypto-tool"),
                         ("gspread", "gspread"), ("google.auth", "google-auth")]:
        try:
            __import__(mod)
            print(f"  [OK ] 패키지 {pipname}")
        except ImportError:
            ok = False
            print(f"  [ERR] 패키지 {pipname} 없음 → pip install {pipname}")
    key = os.path.join(os.path.expanduser("~"), ".secrets", "google-sheets-key.json")
    if os.path.isfile(key):
        print(f"  [OK ] 구글 키 파일 있음 ({key})")
    else:
        ok = False
        print(f"  [ERR] 구글 키 파일 없음: {key}\n        → 기존 PC의 같은 위치 파일을 복사해 오세요(OneDrive·작업 폴더에 두지 말 것).")
    if os.path.isfile(key):
        try:
            sys.path.insert(0, ROOT)
            from sheets_client import open_sheet
            import gsheet_mirror as gm
            for c in COMPANIES:
                t = open_sheet("https://docs.google.com/spreadsheets/d/" + gm.SHEET_IDS[c]).title
                print(f"  [OK ] 구글 시트 접속: {t}")
        except Exception as e:  # noqa: BLE001
            ok = False
            print(f"  [ERR] 구글 시트 접속 실패: {type(e).__name__}: {e}")
    print("\n설치 점검 " + ("완료 — 바로 사용 가능합니다." if ok else "실패 — 위 [ERR] 항목을 해결하세요."))
    return 0 if ok else 1


# ---------------------------------------------------------------- 메인
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--company", choices=COMPANIES)
    ap.add_argument("files", nargs="*")
    ap.add_argument("--no-cafe24", action="store_true", help="플레르: 그 주 자사몰 판매 없음(0 처리)")
    ap.add_argument("--no-sheet", action="store_true", help="구글 시트 입력·대조 생략")
    ap.add_argument("--overwrite", action="store_true", help="시트에 이미 있는 주차도 다시 쓰기")
    ap.add_argument("--check", action="store_true", help="새 PC 설치 점검")
    a = ap.parse_args()
    if a.check:
        sys.exit(check_setup())
    if not a.company or not a.files:
        ap.error("--company 와 원본 파일 경로가 필요합니다")

    company = a.company
    print(f"========== {company} 주간 판매데이터 ==========")
    f = classify(a.files)
    miss = [k for k, n in (("naver", "네이버"), ("coupang", "쿠팡"), ("ohouse", "오늘의집")) if not f[k]]
    if miss:
        die("필수 파일이 없습니다: " + ", ".join({"naver": "네이버(sales_…)", "coupang": "쿠팡(SELLER_INSIGHTS_VENDOR_ITEM_METRICS_…)",
                                                 "ohouse": "오늘의집(주문배송 내역…)"}[k] for k in miss))
    if f["cafe24"] and not HAS_CAFE24[company]:
        die(f"{company}는 자사몰이 없는데 카페24 파일이 들어왔습니다: {os.path.basename(f['cafe24'])} (회사 확인)")
    if HAS_CAFE24[company] and not f["cafe24"] and not a.no_cafe24:
        die("플레르 자사몰(카페24 plaire1054_….csv) 파일이 없습니다.\n"
            "       그 주 자사몰 판매가 없었다면 --no-cafe24 를 붙여 다시 실행하세요(자사몰 0 처리).")

    year, week, s, e = naver_week(f["naver"])
    out = os.path.join(HERE, f"{year}_{company} 주간판매데이터 ({week}).xlsx")
    base = find_base(company, year, week)
    print(f"주차: {week} ({s} ~ {e})")
    for k, n in (("naver", "네이버"), ("coupang", "쿠팡"), ("ohouse", "오늘의집"), ("cafe24", "자사몰")):
        print(f"  {n:5s}: {os.path.basename(f[k]) if f[k] else ('(판매 없음 → 0)' if k == 'cafe24' and HAS_CAFE24[company] else '-')}")
    print(f"  base  : {os.path.basename(base) if base else '(없음 — 이번 주만)'}")

    # 2) 이번 주 집계(주간파일)
    print("\n[1/5] 이번 주 집계")
    import generate_weekly_sales as g
    try:
        g.build(f["naver"], f["coupang"], f["ohouse"], out, week, ohouse_pw=g.DEFAULT_OHOUSE_PW,
                base_path=base, brand=company, cafe24_path=f["cafe24"])
    except PermissionError:
        die(f"파일이 엑셀에 열려 있어 저장하지 못했습니다. 닫고 다시 실행하세요: {os.path.basename(out)}")

    vargs = ["--company", company, "--file", out, "--naver", f["naver"], "--coupang", f["coupang"], "--ohouse", f["ohouse"]]
    if f["cafe24"]:
        vargs += ["--cafe24", f["cafe24"]]

    # 3) 원본 ↔ 주간파일
    print("\n[2/5] 원본 ↔ 이번 주 집계 검증")
    if run("verify_week.py", *vargs, "--no-sheet") != 0:
        die("원본과 집계가 다릅니다. 통합파일·구글 시트에는 쓰지 않았습니다.", 1)

    # 4) 통합파일에 붙이기
    print("\n[3/5] 통합파일에 이번 주 추가")
    import master_book as mb
    prev_master = find_master(company, year, week)
    if not prev_master:
        die(f"통합/ 폴더에 {company} 지난주 통합파일이 없습니다.")
    new_master = os.path.join(HERE, "통합", f"{year}_{company} 주간판매데이터 ({week.replace(' ', '')}).xlsx")
    channels = ["네이버", "쿠팡", "오늘의집"] + (["자사몰"] if HAS_CAFE24[company] else [])
    try:
        mb.append_week(prev_master, out, new_master, channels)
    except PermissionError:
        die(f"통합파일이 엑셀에 열려 있어 저장하지 못했습니다. 닫고 다시 실행하세요: {os.path.basename(new_master)}")
    print(f"  {os.path.basename(prev_master)} + 이번 주 → {os.path.basename(new_master)}")

    # 5) 통합파일 검증
    print("\n[4/5] 통합파일 검증(새 블록·과거 이력·매출데이터·차트)")
    if run("verify_master.py", "--company", company, "--prev", prev_master, "--new", new_master, "--weekly", out) != 0:
        die("통합파일 검증 불일치. 구글 시트에는 쓰지 않았습니다.", 1)

    # 공통 서식(글꼴·크기·테두리·정렬)으로 맞춘 뒤 구글 시트에 올린다 → 엑셀·시트·월간 파일이 같은 모양
    import sheet_style
    for p in (out, new_master):
        sheet_style.normalize_file(p)
    print("  서식 맞춤: 주간파일·통합파일")

    # 6) 구글 시트 교체 + 전 칸 대조
    if not a.no_sheet:
        print("\n[5/5] 구글 시트를 통합파일과 똑같이 교체 + 전 칸 대조")
        if run("gsheet_mirror.py", "--company", company, "--file", new_master) != 0:
            die("구글 시트 대조 불일치.", 1)
    out = new_master

    import openpyxl
    s = openpyxl.load_workbook(out, data_only=True)["매출데이터"]
    w1, w0 = s.cell(3, 2).value, s.cell(3, 3).value
    print(f"\n[요약] 채널 | {w1} | {w0 or '-'}")
    for r in range(4, 10):
        lab = s.cell(r, 1).value
        if lab:
            fmt = lambda v: f"{int(v):,}" if isinstance(v, (int, float)) else "-"
            print(f"  {lab} | {fmt(s.cell(r, 2).value)} | {fmt(s.cell(r, 3).value)}")
    print(f"\n========== 완료: {os.path.basename(out)} ==========")


if __name__ == "__main__":
    main()
