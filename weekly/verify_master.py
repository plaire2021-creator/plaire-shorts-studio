# -*- coding: utf-8 -*-
"""
통합파일(새 양식) 검증
------------------------------------------------------------------
check_new_week : 지난주 통합파일 → 이번 주 통합파일로 한 주 추가했을 때
   [A] 새 블록 = 주간파일 순위시트(칸 단위), 기간·제목·헤더, 병합
   [B] 과거 이력 그대로(순위 탭·기타 탭 전 칸, 병합)
   [C] 매출데이터: B열 = 주간파일 값, 합계 = 채널 합, 기존 열이 정확히 한 칸 이동(수식 계산값·숨김 포함), 차트 범위 +1
   [D] 교차: 새 블록 금액 합 = 매출데이터 채널 값(네이버·오늘의집·자사몰, 쿠팡은 0원 이하 제외분 차이만 허용 보고)
check_all_weeks : 통합파일의 주차별 블록·매출데이터 열이 주간파일들과 같은지(가즈다 새로 만들 때)
"""
import argparse
import sys

import openpyxl

import master_book as mb


class Rep:
    def __init__(self):
        self.secs = []

    def sec(self, t):
        self.secs.append([t, 0, []])

    def ok(self, cond, msg):
        self.secs[-1][1] += 1
        if not cond:
            self.secs[-1][2].append(msg)

    def show(self):
        tot = sum(s[1] for s in self.secs)
        err = sum(len(s[2]) for s in self.secs)
        print("\n================ 통합파일 검증 ================")
        for t, n, e in self.secs:
            print(f"[{'OK ' if not e else 'ERR'}] {t}: {n}건 대조, 불일치 {len(e)}건")
            for m in e[:12]:
                print("       -", m)
            if len(e) > 12:
                print(f"       ... 외 {len(e) - 12}건")
        print(f"----------------------------------------------\n총 {tot}건 대조 / 불일치 {err}건 / 오류율 {(err / tot * 100 if tot else 0):.2f}%")
        return err


def col_hidden(ws, c):
    for d in ws.column_dimensions.values():
        if d.min and d.max and d.min <= c <= d.max:
            return bool(d.hidden)
    return False


def flags(header):
    return [str(h or "").strip() in mb.CODE for h in header]


def block_data(ws, start, w, header):
    f = flags(header)
    last = max((r for r in range(4, ws.max_row + 1)
                if any(ws.cell(r, start + k).value not in (None, "") for k in range(w))), default=3)
    return [[mb.clean(ws.cell(r, start + k).value, f[k]) for k in range(w)] for r in range(4, last + 1)]


def weekly_data(wk, sheet):
    if sheet not in wk.sheetnames:
        return None, []
    ws = wk[sheet]
    ncol = max(c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value not in (None, ""))
    header = [ws.cell(3, c).value for c in range(1, ncol + 1)]
    f = flags(header)
    rows = []
    for r in range(4, ws.max_row + 1):
        row = [mb.clean(ws.cell(r, c).value, f[c - 1]) for c in range(1, ncol + 1)]
        if any(v is not None for v in row):
            rows.append(row)
    return ws.cell(1, 1).value, rows


def weekly_summary(wk):
    s = wk[mb.SUMMARY]
    return {str(s.cell(r, 1).value).strip(): mb.clean(s.cell(r, 2).value) for r in range(4, 12) if s.cell(r, 1).value}


def compare_block(rep, tab, ws, start, w, wk, ch, date_fallback):
    header = [ws.cell(3, start + k).value for k in range(w)]
    dr, rows = weekly_data(wk, ch)
    rep.ok(ws.cell(1, start).value == (dr or date_fallback), f"{tab} 기간: 통합 {ws.cell(1, start).value} / 주간 {dr or date_fallback}")
    got = block_data(ws, start, w, header)
    if not rows:
        rep.ok(all(v is None for r in got for v in r), f"{tab} 판매 없는 주인데 블록에 값 있음")
        return got
    rep.ok(len(got) == len(rows), f"{tab} 행 수: 통합 {len(got)} / 주간 {len(rows)}")
    for i, (a, b) in enumerate(zip(got, rows)):
        rep.ok(a == b, f"{tab} {4 + i}행: 통합 {a} / 주간 {b}")
    return got


def amount_sum(rows, header):
    for name in ("금액", "결제금액"):
        if name in header:
            i = header.index(name)
            return sum(r[i] for r in rows if isinstance(r[i], (int, float)))
    return None


def check_new_week(prev_path, new_path, weekly_path, channels, rep=None):
    rep = rep or Rep()
    prev = openpyxl.load_workbook(prev_path)
    new = openpyxl.load_workbook(new_path)
    wk = openpyxl.load_workbook(weekly_path, data_only=True)
    ws_vals = weekly_summary(wk)
    week = str(wk[mb.SUMMARY].cell(3, 2).value).strip()
    naver_dr = wk["네이버"].cell(1, 1).value

    rep.sec("[A] 새 주차 블록 ↔ 주간파일")
    sums = {}
    for ch in channels:
        tab = mb.TAB[ch]
        p, n = prev[tab], new[tab]
        sp, sn = mb.block_starts(p), mb.block_starts(n)
        rep.ok(sn[:-1] == sp and len(sn) == len(sp) + 1, f"{tab} 블록 위치: 기존 {len(sp)}개 → {len(sn)}개")
        start = sn[-1]
        rep.ok(start - sp[-1] == (sp[-1] - sp[-2] if len(sp) > 1 else start - sp[-1]), f"{tab} 블록 간격이 직전과 다름")
        w = mb.block_width(p, sp[-1])
        rep.ok(n.cell(2, start).value == p.cell(2, sp[-1]).value, f"{tab} 제목: {n.cell(2, start).value}")
        hdr_new = [n.cell(3, start + k).value for k in range(w)]
        rep.ok(hdr_new == [p.cell(3, sp[-1] + k).value for k in range(w)], f"{tab} 헤더 다름: {hdr_new}")
        for r in (1, 2):
            rep.ok(any(m.min_row == r and m.min_col == start and m.max_col == start + w - 1 for m in n.merged_cells.ranges),
                   f"{tab} {r}행 병합 없음")
        got = compare_block(rep, tab, n, start, w, wk, ch, naver_dr)
        sums[ch] = amount_sum(got, hdr_new)

    rep.sec("[B] 과거 이력 보존(기존 칸 전부)")
    for ws in prev.worksheets:
        if ws.title == mb.SUMMARY:
            continue
        n = new[ws.title]
        diff = [ws.cell(r, c).coordinate for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)
                if ws.cell(r, c).value != n.cell(r, c).value]
        rep.ok(not diff, f"{ws.title}: 기존 칸 {len(diff)}개가 바뀜 {diff[:5]}")
        lost = {str(m) for m in ws.merged_cells.ranges} - {str(m) for m in n.merged_cells.ranges}
        rep.ok(not lost, f"{ws.title}: 병합 사라짐 {sorted(lost)[:5]}")

    rep.sec("[C] 매출데이터")
    p, n = prev[mb.SUMMARY], new[mb.SUMMARY]
    labels, last = mb.summary_layout(p)
    rep.ok(n.cell(3, 2).value == week, f"B3 = {n.cell(3, 2).value} (기대 {week})")
    chan = [k for k in labels if k != "합계"]
    for lab in chan:
        want = ws_vals.get(lab, 0) or 0
        rep.ok(n.cell(labels[lab], 2).value == want, f"B {lab}: {n.cell(labels[lab], 2).value} / 주간 {want}")
    rep.ok(n.cell(labels["합계"], 2).value == sum(n.cell(labels[k], 2).value or 0 for k in chan), "B 합계 ≠ 채널 합")
    rep.ok(n.cell(labels["합계"], 2).value == ws_vals.get("합계"), f"B 합계 {n.cell(labels['합계'], 2).value} / 주간 {ws_vals.get('합계')}")
    for c in range(2, last + 1):
        for r in range(1, p.max_row + 1):
            a, b = mb.eval_cell(p, p.cell(r, c).value), mb.eval_cell(n, n.cell(r, c + 1).value)
            same = a == b or (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 0.5)
            rep.ok(same, f"{p.cell(r, c).coordinate}→{n.cell(r, c + 1).coordinate}: {a} / {b}")
        rep.ok(col_hidden(p, c) == col_hidden(n, c + 1), f"숨김 {p.cell(1, c).column_letter}→{n.cell(1, c + 1).column_letter}")
    pf = [s.val.numRef.f for ch in p._charts for s in ch.series]
    nf = [s.val.numRef.f for ch in n._charts for s in ch.series]
    rep.ok(len(pf) == len(nf), f"차트 계열 수 {len(pf)}→{len(nf)}")
    for a, b in zip(pf, nf):
        exp = mb.re.sub(r"\$([A-Z]+)\$(\d+)$", lambda m: f"${mb.L(openpyxl.utils.column_index_from_string(m.group(1)) + 1)}${m.group(2)}", a)
        rep.ok(b == exp, f"차트 범위 {a} → {b} (기대 {exp})")

    rep.sec("[D] 교차: 새 블록 금액 합 = 매출데이터")
    for ch in channels:
        s = sums.get(ch)
        v = n.cell(labels[ch], 2).value if ch in labels else None
        if ch == "쿠팡":
            rep.ok(s is not None and v is not None and s >= v, f"쿠팡 블록 합 {s} < 매출 {v}")
        else:
            rep.ok(s == v or (not s and not v), f"{ch} 블록 합 {s} / 매출데이터 {v}")
    return rep


def check_all_weeks(master_path, weekly_paths, channels, rep=None):
    """weekly_paths: 오래된 주 → 최신 주 순서. 통합파일 블록 i번째 = 주간파일 i번째."""
    rep = rep or Rep()
    m = openpyxl.load_workbook(master_path)
    rep.sec(f"전체 주차({len(weekly_paths)}주) 블록 ↔ 주간파일")
    for ch in channels:
        ws = m[mb.TAB[ch]]
        starts = mb.block_starts(ws)
        rep.ok(len(starts) == len(weekly_paths), f"{ws.title} 블록 {len(starts)}개 / 주간파일 {len(weekly_paths)}개")
        for s0, wp in zip(starts, weekly_paths):
            wk = openpyxl.load_workbook(wp, data_only=True)
            w = mb.block_width(ws, s0)
            got = compare_block(rep, f"{ws.title}[{wk[mb.SUMMARY].cell(3, 2).value}]", ws, s0, w, wk, ch, wk["네이버"].cell(1, 1).value)
    rep.sec("전체 주차 매출데이터 ↔ 주간파일")
    s = m[mb.SUMMARY]
    labels, last = mb.summary_layout(s)
    rep.ok(last - 1 == len(weekly_paths), f"매출데이터 주차 열 {last - 1}개 / 주간파일 {len(weekly_paths)}개")
    for i, wp in enumerate(reversed(weekly_paths)):        # B열 = 최신
        c = 2 + i
        wk = openpyxl.load_workbook(wp, data_only=True)
        v = weekly_summary(wk)
        rep.ok(s.cell(3, c).value == wk[mb.SUMMARY].cell(3, 2).value, f"{s.cell(3, c).coordinate} 주차 {s.cell(3, c).value}")
        for lab, r in labels.items():
            want = v.get(lab, 0) or 0
            got = mb.eval_cell(s, s.cell(r, c).value)
            rep.ok(got == want, f"{s.cell(r, c).coordinate} {lab}: 통합 {got} / 주간 {want}")
    return rep


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--prev", required=True)
    ap.add_argument("--new", required=True)
    ap.add_argument("--weekly", required=True)
    ap.add_argument("--company", required=True, choices=["플레르가구", "가즈다가구"])
    a = ap.parse_args()
    ch = ["네이버", "쿠팡", "오늘의집"] + (["자사몰"] if a.company == "플레르가구" else [])
    sys.exit(1 if check_new_week(a.prev, a.new, a.weekly, ch).show() else 0)


if __name__ == "__main__":
    main()
