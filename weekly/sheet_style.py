# -*- coding: utf-8 -*-
"""
매출 엑셀 공통 서식 — 주간파일·통합파일(=구글 시트)·월간 매출통계·토탈을 모두 같은 모양으로 맞춘다.
(weekly/ 와 monthly/ 에 같은 파일이 하나씩 있다. 고칠 때는 둘 다 고칠 것)

값·수식·병합·열 너비·차트는 건드리지 않고, 글꼴·크기·테두리·정렬·줄 높이·머리글/강조 색을 바꾼다.
표시 형식은 '일반'인 숫자 칸만 천 단위 쉼표(#,##0)로 (코드·순위·%·날짜 형식은 그대로).

표 찾는 법(양식이 조금씩 달라도 되게 칸 위치를 고정하지 않는다):
  - 시트를 '빈 열'로 나눈 덩어리(블록) 하나하나를, 다시 '빈 행'으로 나눈다.
  - 덩어리 맨 위에서 칸이 하나뿐인 줄(병합 포함) = 제목. 마지막 제목 줄은 큰 제목, 그 위는 날짜 같은 작은 제목.
  - 그 아래부터 표. 숫자가 하나도 없는 첫 줄들(최대 3줄) = 머리글, 나머지 = 본문.
  - 본문 첫 열이 짧은 글자(12자 이하)이고 오른쪽에 숫자가 있으면 = 행 머리글(채널·몰 이름).

공통 서식:
  글꼴 맑은 고딕 / 큰 제목 14 굵게, 작은 제목 11 굵게 / 머리글 11 굵게 + 연한 파랑 / 본문 11
  테두리 표 안 모든 칸 가는 선(제목은 선 없음) / 세로 가운데
  가로: 머리글·행 머리글·순위·코드 가운데, 숫자 오른쪽, 글자 왼쪽
  줄 높이: 큰 제목 30, 작은 제목 20, 머리글 24, 본문 21
  노란 계열 강조색(최신 주차 등)은 모두 같은 연노랑으로
"""
import copy
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "맑은 고딕"
SIZE_TITLE, SIZE_LABEL, SIZE = 14, 11, 11
H_TITLE, H_LABEL, H_HEAD, H_BODY = 30, 20, 24, 21
THIN = Side(style="thin", color="FFBFBFBF")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NONE = Border()
HEAD_FILL = PatternFill("solid", fgColor="FFDDE5F3")
HIGHLIGHT = PatternFill("solid", fgColor="FFFFF2CC")
CODE_WORDS = ("코드", "ID", "id", "번호")


def _num(v):
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    return isinstance(v, str) and v.startswith("=")          # 수식 = 숫자 칸


def _yellow(fill):
    if not fill or fill.fill_type != "solid":
        return False
    rgb = str(fill.fgColor.rgb or "")[-6:].upper()
    if not re.fullmatch(r"[0-9A-F]{6}", rgb):
        return False
    r, g, b = int(rgb[0:2], 16), int(rgb[2:4], 16), int(rgb[4:6], 16)
    return r >= 0xF0 and g >= 0xC8 and b <= 0xD8 and r - b >= 0x28   # 노랑~연노랑


def _font(c, size, bold):
    f = copy.copy(c.font)
    c.font = Font(name=FONT, sz=size, b=bold, i=f.i, color=f.color, u=f.u, strike=f.strike)


def _align(c, h):
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=False, indent=1 if h == "left" else 0)


def normalize_sheet(ws):
    maxr, maxc = ws.max_row, ws.max_column
    top = {}            # 병합 안 칸 → 병합 왼쪽 위 칸
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                top[(r, c)] = (m.min_row, m.min_col)

    vals = {}
    for row in ws.iter_rows(min_row=1, max_row=maxr, max_col=maxc):
        for x in row:
            if x.value is not None and str(x.value).strip() != "":
                vals[(x.row, x.column)] = x.value

    def val(r, c):
        return vals.get(top.get((r, c), (r, c)))

    # 블록 사이 좁은 빈 열에 실수로 들어간 글자 하나('#' 등) 때문에 블록이 붙지 않게: 칸 하나뿐인 좁은 열은 빈 열로 본다
    from openpyxl.utils import get_column_letter
    count = {}
    for (r, c) in vals:
        count[c] = count.get(c, 0) + 1
    for (r, c), t in top.items():
        if t in vals:
            count[c] = max(count.get(c, 0), 2)

    def gap(c):
        w = ws.column_dimensions[get_column_letter(c)].width
        return count.get(c, 0) <= 1 and w is not None and w < 4

    used_cols = sorted(c for c in count if not gap(c))
    blocks, cur = [], []
    for c in used_cols:
        if cur and c != cur[-1] + 1:
            blocks.append(cur)
            cur = []
        cur.append(c)
    if cur:
        blocks.append(cur)

    heights = {}
    for cols in blocks:
        c0, c1 = cols[0], cols[-1]
        rows = sorted({r for r in range(1, maxr + 1) if any(val(r, c) is not None for c in cols)})
        segs, s = [], []
        for r in rows:
            if s and r != s[-1] + 1:
                segs.append(s)
                s = []
            s.append(r)
        if s:
            segs.append(s)
        for seg in segs:
            def units(r):
                return {top.get((r, c), (r, c)) for c in cols if val(r, c) is not None}

            # 1) 제목 줄
            i = 0
            while i < len(seg) and len(units(seg[i])) <= 1 and not any(_num(val(seg[i], c)) for c in cols):
                i += 1
            titles, body = seg[:i], seg[i:]
            if body and len(titles) == len(seg):
                titles, body = seg, []
            for k, r in enumerate(titles):
                big = k == len(titles) - 1
                for c in cols:
                    x = ws.cell(r, c)
                    x.border = NONE
                    if val(r, c) is not None and top.get((r, c), (r, c)) == (r, c):
                        _font(x, SIZE_TITLE if big else SIZE_LABEL, True)
                        x.alignment = Alignment(horizontal=x.alignment.horizontal, vertical="center")
                        if _yellow(x.fill):
                            x.fill = HIGHLIGHT
                heights[r] = max(heights.get(r, 0), H_TITLE if big else H_LABEL)
            if not body:
                continue

            # 2) 머리글 줄 (숫자 없는 첫 줄들, 아래에 숫자 줄이 있을 때만)
            nh = 0
            while nh < min(3, len(body) - 1) and not any(_num(val(body[nh], c)) for c in cols) \
                    and any(_num(val(r, c)) for r in body[nh + 1:] for c in cols):
                nh += 1
            heads, data = body[:nh], body[nh:]
            names = {c: " ".join(str(val(r, c) or "") for r in heads) for c in cols}

            # 열 성격: 코드(가운데) · 순위(첫 열 작은 정수, 가운데) · 행 머리글(첫 열 글자 + 오른쪽 숫자)
            def col_nums(c):
                return [val(r, c) for r in data if _num(val(r, c))]
            code_col = {c for c in cols if any(w in names[c] for w in CODE_WORDS)}
            for c in cols:
                ns = [v for v in col_nums(c) if isinstance(v, (int, float))]
                if ns and len(ns) == len(col_nums(c)) and all(float(v).is_integer() and abs(v) >= 1e9 for v in ns):
                    code_col.add(c)          # 9자리 넘는 정수만 있는 열 = 상품 번호
            first = c0
            fn = [val(r, first) for r in data if val(r, first) is not None]
            rank_col = bool(fn) and all(isinstance(v, (int, float)) and float(v).is_integer() and 0 < v < 1000 for v in fn)
            # 행 머리글은 채널·몰 이름처럼 짧은 이름만 (긴 상품명 열은 보통 글자처럼 왼쪽)
            row_head = bool(fn) and all(not _num(v) and len(str(v)) <= 12 for v in fn) and any(col_nums(c) for c in cols[1:])

            for r in heads:
                for c in cols:
                    x = ws.cell(r, c)
                    x.border = GRID
                    _font(x, SIZE, True)
                    _align(x, "center")
                    x.fill = HIGHLIGHT if _yellow(x.fill) else HEAD_FILL
                heights[r] = max(heights.get(r, 0), H_HEAD)
            for r in data:
                for c in cols:
                    x = ws.cell(r, c)
                    x.border = GRID
                    v = val(r, c)
                    if _yellow(x.fill):
                        x.fill = HIGHLIGHT
                    if c == first and row_head:
                        _font(x, SIZE, True)
                        _align(x, "center")
                    else:
                        _font(x, SIZE, False)
                        if (c == first and rank_col) or c in code_col:
                            _align(x, "center")
                        elif _num(v):
                            _align(x, "right")
                            if x.number_format == "General":       # 금액·수량은 천 단위 쉼표 (코드·순위·% 는 그대로)
                                x.number_format = "#,##0"
                        else:
                            _align(x, "left")
                heights[r] = max(heights.get(r, 0), H_BODY)

    for r, h in heights.items():
        ws.row_dimensions[r].height = h


def normalize(wb):
    for ws in wb.worksheets:
        normalize_sheet(ws)
    return wb


def normalize_file(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    normalize(wb)
    wb.save(path)


if __name__ == "__main__":
    import sys
    for p in sys.argv[1:]:
        normalize_file(p)
        print("서식 맞춤:", p)
