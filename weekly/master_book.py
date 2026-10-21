# -*- coding: utf-8 -*-
"""
새 양식(통합파일) 엔진 — 2026-10 사용자 확정 양식
------------------------------------------------------------------
통합파일 1개에 전체 이력이 쌓인다(플레르: 2024.07~ / 가즈다: 2026.08~).

  순위 탭(네이버TOP20·쿠팡TOP20·오집TOP20·[자사몰TOP20])
    주차 블록이 왼쪽(과거)→오른쪽(최신). 1행 기간(병합) / 2행 'OO 판매순위'(노랑, 병합) /
    3행 헤더 / 4행~ 데이터, 블록 사이 빈 열 1개. 새 블록 서식은 직전(마지막) 블록을 그대로 따른다.
  매출데이터
    최신 주차가 B열(노랑), 과거 주차는 오른쪽으로 전부 유지. 새 주차를 넣을 때 기존 열을 한 칸씩
    오른쪽으로 옮기고(수식은 위치에 맞게 변환) 차트 범위 끝을 한 칸 늘린다.

데이터는 generate_weekly_sales.py 가 만든 주간파일(구양식 xlsx)의 순위 시트·매출데이터에서 그대로 옮긴다.
"""
import copy
import re

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

PRICE = {"금액", "결제금액", "총합계"}
CODE = {"상품 코드", "상품코드", "옵션코드", "상품아이디"}
TAB = {"네이버": "네이버TOP20", "쿠팡": "쿠팡TOP20", "오늘의집": "오집TOP20", "자사몰": "자사몰TOP20"}
TITLE = {"네이버": "네이버 판매순위", "쿠팡": "쿠팡 판매순위", "오늘의집": "오늘의집 판매순위", "자사몰": "자사몰 판매순위"}
SUMMARY = "매출데이터"


# ---------------------------------------------------------------- 공통
def clean(v, code=False):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return None
        if code and re.fullmatch(r"[1-9]\d{0,14}", s):
            return int(s)
    return v


def style_of(cell):
    """워크북에 묶이지 않은 스타일 사본(다른 파일로도 옮길 수 있게)."""
    return {"font": copy.copy(cell.font), "border": copy.copy(cell.border), "fill": copy.copy(cell.fill),
            "alignment": copy.copy(cell.alignment), "number_format": cell.number_format,
            "protection": copy.copy(cell.protection)}


def apply_style(cell, st):
    cell.font, cell.border, cell.fill = st["font"], st["border"], st["fill"]
    cell.alignment, cell.number_format, cell.protection = st["alignment"], st["number_format"], st["protection"]


# ---------------------------------------------------------------- 열 너비(묶음 설정 분리)
def _new_dim(ws, d, mn, mx):
    from openpyxl.worksheet.dimensions import ColumnDimension
    nd = ColumnDimension(ws, index=L(mn), width=d.width, bestFit=d.bestFit, hidden=d.hidden,
                         outlineLevel=d.outlineLevel, collapsed=d.collapsed)
    nd.min, nd.max = mn, mx
    nd._style = copy.copy(d._style)
    ws.column_dimensions[L(mn)] = nd


def split_dims(ws, lo, hi):
    """lo~hi 열에 걸친 '여러 열 묶음' 너비 설정을 열 단위로 나눈다(그대로 두고 한 열만 고치면 파일이 깨짐)."""
    for key, d in list(ws.column_dimensions.items()):
        mn, mx = d.min, d.max
        if mn is None or mx is None or mn == mx or mx < lo or mn > hi:
            continue
        del ws.column_dimensions[key]
        if mn < lo:
            _new_dim(ws, d, mn, lo - 1)
        for c in range(max(mn, lo), min(mx, hi) + 1):
            _new_dim(ws, d, c, c)
        if mx > hi:
            _new_dim(ws, d, hi + 1, mx)


# ---------------------------------------------------------------- 순위 탭
def block_starts(ws):
    s = {c for r in (1, 2) for c in range(1, ws.max_column + 1) if ws.cell(r, c).value not in (None, "")}
    return sorted(s)


def block_width(ws, start):
    k = 1
    while ws.cell(3, start + k).value not in (None, ""):
        k += 1
    return k                                   # 순위열(헤더 빈칸) 포함 폭


def capture_profile(ws, start=None):
    """마지막 블록(또는 start 블록)의 서식·폭·헤더를 그대로 떠 둔다."""
    starts = block_starts(ws)
    start = start or starts[-1]
    w = block_width(ws, start)
    prev = [s for s in starts if s < start]
    gap_col = start - 1
    merged = {r: any(m.min_row == r and m.min_col == start and m.max_col == start + w - 1
                     for m in ws.merged_cells.ranges) for r in (1, 2)}
    last_row = max((r for r in range(4, ws.max_row + 1)
                    if any(ws.cell(r, start + k).value not in (None, "") for k in range(w))), default=4)
    data, empty = [], []
    for k in range(w):
        filled = [ws.cell(r, start + k) for r in range(4, last_row + 1) if ws.cell(r, start + k).value not in (None, "")]
        blanks = [ws.cell(r, start + k) for r in range(4, last_row + 1) if ws.cell(r, start + k).value in (None, "")]
        data.append(style_of(filled[0] if filled else ws.cell(4, start + k)))
        empty.append(style_of(blanks[0] if blanks else ws.cell(4, start + k)))
    return {
        "width": w,
        "stride": (start - prev[-1]) if prev else w + 1,
        "title": ws.cell(2, start).value,
        "header": [ws.cell(3, start + k).value for k in range(w)],
        "rows": {r: [style_of(ws.cell(r, start + k)) for k in range(w)] for r in (1, 2, 3)},
        "data": data, "empty": empty, "merged": merged,
        "widths": [ws.column_dimensions[L(start + k)].width for k in range(w)],
        "gap_width": ws.column_dimensions[L(gap_col)].width if gap_col >= 1 else None,
        "heights": {r: ws.row_dimensions[r].height for r in (1, 2, 3)},
        "data_height": ws.row_dimensions[4].height,
    }


def write_block(ws, start, prof, date_range, rows):
    """rows: 주간파일 순위시트 4행~ 값 목록(순위열 포함). 헤더/제목은 prof(직전 블록) 그대로."""
    w = prof["width"]
    split_dims(ws, max(1, start - 1), start + w)
    if start > 1 and prof["gap_width"]:
        ws.column_dimensions[L(start - 1)].width = prof["gap_width"]
    for k in range(w):
        if prof["widths"][k]:
            ws.column_dimensions[L(start + k)].width = prof["widths"][k]
    for r, val in ((1, date_range), (2, prof["title"]), (3, None)):
        for k in range(w):
            c = ws.cell(r, start + k)
            apply_style(c, prof["rows"][r][k])
            c.value = prof["header"][k] if r == 3 else (val if k == 0 else None)
        if r in (1, 2) and prof["merged"][r]:
            ws.merge_cells(start_row=r, start_column=start, end_row=r, end_column=start + w - 1)
        if prof["heights"][r] and ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = prof["heights"][r]
    header = prof["header"]
    body = rows if rows else [[None] * w]          # 판매 없는 주: 테두리만 있는 빈 줄 1개
    for i, row in enumerate(body):
        r = 4 + i
        for k in range(w):
            v = clean(row[k] if k < len(row) else None, code=str(header[k] or "").strip() in CODE)
            c = ws.cell(r, start + k)
            c.value = v
            apply_style(c, prof["data"][k] if v is not None else prof["empty"][k])
            if v is not None and str(header[k] or "").strip() in PRICE:
                c.number_format = "#,##0"
            if v is not None and str(header[k] or "").strip() == "옵션":
                c.alignment = Alignment(horizontal="left", vertical=c.alignment.vertical, wrap_text=c.alignment.wrap_text)
        if prof["data_height"] and ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = prof["data_height"]
    return len(rows)


def weekly_rank_rows(wb, sheet):
    """주간파일(구양식) 순위 시트 → (기간, 행 목록)."""
    if sheet not in wb.sheetnames:
        return None, []
    ws = wb[sheet]
    ncol = max(c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value not in (None, ""))
    rows = []
    for r in range(4, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ncol + 1)]
        if any(v not in (None, "") for v in vals):
            rows.append(vals)
    return ws.cell(1, 1).value, rows


# ---------------------------------------------------------------- 매출데이터
def eval_cell(ws, v):
    """매출데이터 숫자 칸 값 계산(수식 =SUM(X4:X7), =숫자-숫자 지원). 못 하면 None."""
    if not (isinstance(v, str) and v.startswith("=")):
        return v
    f = v[1:].replace(" ", "")
    m = re.fullmatch(r"SUM\(([^()]*)\)", f, re.I)
    if m:                                       # SUM(A1:B3), SUM(A1), SUM(A1,B2:B4,100)
        from openpyxl.utils import column_index_from_string as CI
        tot = 0
        for part in m.group(1).split(","):
            pm = re.fullmatch(r"\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?", part, re.I)
            if pm:
                c1, r1, c2, r2 = pm.groups()
                c2, r2 = c2 or c1, r2 or r1
                for c in range(CI(c1.upper()), CI(c2.upper()) + 1):
                    for r in range(int(r1), int(r2) + 1):
                        x = eval_cell(ws, ws.cell(r, c).value)
                        if isinstance(x, (int, float)):
                            tot += x
                        elif x not in (None, ""):
                            return None
            elif re.fullmatch(r"-?\d+(\.\d+)?", part):
                tot += float(part)
            else:
                return None
        return tot
    if re.fullmatch(r"[\d.+\-*/()]+", f):
        return eval(f)  # noqa: S307  숫자·사칙연산만 허용(위 정규식)
    return None


def summary_layout(ws):
    labels = {}
    for r in range(4, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v not in (None, "") and str(v).strip():
            labels[str(v).strip()] = r
    weeks = [c for c in range(2, ws.max_column + 1) if ws.cell(3, c).value not in (None, "")]
    return labels, (max(weeks) if weeks else 1)


def summary_insert_week(ws, week, values):
    """기존 주차 열을 한 칸씩 오른쪽으로 옮기고 B열에 새 주차. values: {'네이버':…, …, '합계':…}."""
    labels, last = summary_layout(ws)
    maxr = ws.max_row
    split_dims(ws, 2, last + 1)
    for c in range(last, 1, -1):                       # 오른쪽부터 한 칸씩 이동
        for r in range(1, maxr + 1):
            src, dst = ws.cell(r, c), ws.cell(r, c + 1)
            v = src.value
            if isinstance(v, str) and v.startswith("="):
                v = Translator(v, origin=src.coordinate).translate_formula(dst.coordinate)
            dst.value = v
            dst._style = copy.copy(src._style)
        cd_s, cd_d = ws.column_dimensions[L(c)], ws.column_dimensions[L(c + 1)]
        cd_d.width, cd_d.hidden = cd_s.width, cd_s.hidden
    # 새 B열: 옛 B(최신 서식, 이제 C) 서식 복사 → 값 기록
    for r in range(1, maxr + 1):
        ws.cell(r, 2)._style = copy.copy(ws.cell(r, 3)._style)
        ws.cell(r, 2).value = None
    ws.column_dimensions["B"].hidden = False
    ws.cell(3, 2).value = week
    for lab, r in labels.items():
        if lab in values:
            ws.cell(r, 2).value = values[lab]
    # 옛 최신 열(C)은 일반 주차 서식(옛 C = 지금 D)으로
    if last >= 3:
        for r in range(3, max(labels.values()) + 1):
            ws.cell(r, 3)._style = copy.copy(ws.cell(r, 4)._style)
    for col in (2, 3):
        for r in range(4, max(labels.values()) + 1):
            ws.cell(r, col).number_format = "#,##0"
    extend_chart_refs(ws, 1)


def extend_chart_refs(ws, n):
    """차트 계열 범위 '$B$r:$XX$r' 의 끝 열을 n칸 늘림(새 주차 포함)."""
    from openpyxl.utils import column_index_from_string as CI

    def fix(f):
        return re.sub(r"\$B\$(\d+):\$([A-Z]+)\$(\d+)",
                      lambda m: f"$B${m.group(1)}:${L(CI(m.group(2)) + n)}${m.group(3)}", f)
    for ch in ws._charts:
        for s in ch.series:
            for ref in (getattr(s.cat, "numRef", None), getattr(s.cat, "strRef", None), getattr(s.val, "numRef", None)):
                if ref is not None and ref.f:
                    ref.f = fix(ref.f)


# ---------------------------------------------------------------- 주차 추가
def append_week(master_path, weekly_path, out_path, channels, profiles=None):
    """통합파일(master_path)에 주간파일(weekly_path)의 이번 주를 붙여 out_path로 저장.
    channels: 이 회사의 채널 목록(가즈다는 자사몰 제외)."""
    wb = openpyxl.load_workbook(master_path)
    wk = openpyxl.load_workbook(weekly_path, data_only=True)
    res = append_week_wb(wb, wk, channels, profiles)
    wb.save(out_path)
    return res


def append_week_wb(wb, wk, channels, profiles=None):
    week = str(wk[SUMMARY].cell(3, 2).value).strip()
    added = {}
    for ch in channels:
        ws = wb[TAB[ch]]
        prof = (profiles or {}).get(ch) or capture_profile(ws)
        starts = block_starts(ws)
        start = (starts[-1] + prof["stride"]) if starts else 1
        date_range, rows = weekly_rank_rows(wk, ch)
        if date_range is None:                         # 그 주 순위 시트 없음(자사몰 판매 0) → 기간은 네이버에서
            date_range = wk["네이버"].cell(1, 1).value
        added[ch] = (start, write_block(ws, start, prof, date_range, rows))
    src = wk[SUMMARY]
    vals = {str(src.cell(r, 1).value).strip(): clean(src.cell(r, 2).value)
            for r in range(4, 12) if src.cell(r, 1).value}
    labels, _ = summary_layout(wb[SUMMARY])
    vals = {k: (vals.get(k, 0) if k != "합계" else None) for k in labels}
    vals["합계"] = sum(v for k, v in vals.items() if k != "합계" and isinstance(v, (int, float)))
    summary_insert_week(wb[SUMMARY], week, vals)
    return week, added, vals


# ---------------------------------------------------------------- 새 통합파일 만들기(가즈다 등)
def build_master(company, weekly_paths, style_src_path, out_path, channels):
    """주간파일들(오래된 주→최신 주)로 새 통합파일을 만든다. 서식은 style_src(플레르 통합파일)를 그대로 따른다."""
    from openpyxl.chart import LineChart, Reference
    src = openpyxl.load_workbook(style_src_path)
    profiles = {ch: capture_profile(src[TAB[ch]]) for ch in channels}
    ss = src[SUMMARY]
    s_labels, _ = summary_layout(ss)                      # 플레르: 네이버4 쿠팡5 오늘의집6 자사몰7 합계8
    labels = [ch for ch in ("네이버", "쿠팡", "오늘의집", "자사몰") if ch in channels] + ["합계"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in [SUMMARY, "네이버TOP20", "쿠팡TOP20", "오집TOP20"] + (["자사몰TOP20"] if "자사몰" in channels else []):
        wb.create_sheet(name)
    sm = wb[SUMMARY]
    # 매출데이터 뼈대: 원본 행 r ↔ 새 행 매핑(자사몰 행이 없으면 합계 이하가 한 칸 위)
    rowmap = {1: 1, 2: 2, 3: 3}
    for i, lab in enumerate(labels):
        rowmap[4 + i] = s_labels[lab]
    extra = ss.max_row - s_labels["합계"]
    for k in range(1, extra + 1):
        rowmap[4 + len(labels) - 1 + k] = s_labels["합계"] + k
    for r_new, r_src in rowmap.items():
        apply_style(sm.cell(r_new, 1), style_of(ss.cell(r_src, 1)))
        if ss.row_dimensions[r_src].height:
            sm.row_dimensions[r_new].height = ss.row_dimensions[r_src].height
    sm.cell(1, 1).value = f"{company} 주간판매 데이터"
    for i, lab in enumerate(labels):
        sm.cell(4 + i, 1).value = lab
    sm.column_dimensions["A"].width = ss.column_dimensions["A"].width

    for wp in weekly_paths:
        wk = openpyxl.load_workbook(wp, data_only=True)
        append_week_wb(wb, wk, channels, profiles)          # 모든 주차를 플레르 서식·너비·간격으로

    # 매출데이터 서식: B = 최신(노랑), C~ = 일반 — 플레르 B/C열 서식 그대로
    _, last = summary_layout(sm)
    for c in range(2, last + 1):
        src_col = 2 if c == 2 else 3
        for r_new, r_src in rowmap.items():
            v = sm.cell(r_new, c).value
            apply_style(sm.cell(r_new, c), style_of(ss.cell(r_src, src_col)))
            if 4 <= r_new < 4 + len(labels):
                sm.cell(r_new, c).number_format = "#,##0"
        sm.column_dimensions[L(c)].width = ss.column_dimensions["B" if c == 2 else "C"].width

    # 차트: 플레르와 같은 꺾은선 2개(합계 / 채널별), 범위 B열~마지막 주차
    def chart(title, rows):
        ch = LineChart()
        ch.title = title
        for r in rows:
            ch.add_data(Reference(sm, min_col=1, max_col=last, min_row=r, max_row=r), titles_from_data=True, from_rows=True)
        ch.set_categories(Reference(sm, min_col=2, max_col=last, min_row=3, max_row=3))
        ch.width, ch.height = 32, 11
        return ch
    tot_row = 4 + len(labels) - 1
    sm.add_chart(chart("합계", [tot_row]), "A10")
    c2 = chart(f"{company} 주간매출", list(range(4, tot_row)))
    c2.height = 18
    sm.add_chart(c2, "A34")
    wb.save(out_path)
    return out_path
