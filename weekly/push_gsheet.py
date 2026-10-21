# -*- coding: utf-8 -*-
"""
주간판매데이터 xlsx → 구글 시트('플레르가구 주간 판매 데이터' / '가즈다가구 주간 판매 데이터') 입력
------------------------------------------------------------------
시트 양식(사장님이 만든 양식을 그대로 따름):
  매출데이터 : 3행=주차(B열부터 오른쪽으로 누적), A열=채널(네이버/쿠팡/오늘의집/[자사몰]/합계).
               최신 주차 열만 연노랑, 나머지는 흰색.
  순위 탭    : 3행=주차 라벨(블록 폭만큼 병합) / 4행=헤더 / 5행~=데이터, 블록 사이 빈 열 1개.
               블록 = 주간파일 순위 시트의 3행(헤더)~끝을 그대로 옮긴 것(A열=순위).

규칙:
  - 이미 있는 주차는 건드리지 않는다(사장님이 손으로 다듬은 값 보호). 덮어쓰려면 --overwrite.
  - 새 블록의 3~4행(라벨·헤더 서식·병합)은 직전 블록을 복사하고, 헤더 이름도 시트 것을 유지.
  - 그 주에 순위 시트가 없으면(예: 자사몰 판매 0건) 라벨·헤더만 있는 빈 블록을 넣는다.

실행:
  python push_gsheet.py --company 가즈다가구 --file "2026_가즈다가구 주간판매데이터 (10월 12주차).xlsx"
  python push_gsheet.py --company 가즈다가구 --file ... --dry-run     # 쓰지 않고 계획만 출력
"""
import argparse
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from sheets_client import open_sheet  # noqa: E402  (PLAIRE AI\sheets_client.py)

SHEET_IDS = {
    "플레르가구": "1caviDCetj3RsdhRmgqRarfK7y6ee8Kt7t9mDe08VmU4",
    "가즈다가구": "1VoCNP7TJ01JTa0uR8XDM1jcuumyU4ts6A0yrXfJhXgo",
}
RANK_TABS = ["네이버", "쿠팡", "오늘의집", "자사몰"]
LATEST = {"red": 1.0, "green": 0.945, "blue": 0.8}     # 연노랑 (fff1cc)
WHITE = {"red": 1.0, "green": 1.0, "blue": 1.0}


def week_key(label):
    m = re.search(r"(\d+)\s*월\s*0*(\d+)\s*주차", str(label))
    return (int(m.group(1)), int(m.group(2))) if m else None


def week_from_file(path):
    m = re.search(r"\(([^)]*주차)\)", os.path.basename(path))
    return m.group(1).strip() if m else None


def gs_range(r0, c0, r1, c1):
    """0-기준 행/열 → GridRange용 dict 조각."""
    return {"startRowIndex": r0, "endRowIndex": r1, "startColumnIndex": c0, "endColumnIndex": c1}


CODE_HEADERS = {"상품 코드", "상품코드", "옵션코드", "상품아이디"}


def clean(v, code=False):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if code and isinstance(v, str) and re.fullmatch(r"[1-9]\d{0,14}", v.strip()):
        return int(v.strip())          # 코드 열만 숫자로(기존 시트처럼 오른쪽 정렬). 옵션 '2100' 등은 글자 유지
    return v


def push_summary(sh, ws, wb, week, overwrite, dry, reqs):
    vals = ws.get_all_values()
    row3 = vals[2] if len(vals) > 2 else []
    cols = [i for i, v in enumerate(row3) if i > 0 and v.strip()]
    hit = [i for i in cols if week_key(row3[i]) == week_key(week)]
    if hit and not overwrite:
        print(f"  [매출데이터] {week} 이미 있음 → 건너뜀")
        return
    col = hit[0] if hit else (max(cols) + 1 if cols else 1)
    labels = {str(r[0]).strip(): i for i, r in enumerate(vals) if i >= 3 and r and str(r[0]).strip()}
    last_row = max(labels.values()) if labels else 7

    src = wb["매출데이터"]
    weekly = {}
    for r in range(4, 12):
        lab = str(src.cell(r, 1).value or "").strip()
        if lab:
            weekly[lab] = clean(src.cell(r, 2).value)

    cells = []
    for lab, ri in labels.items():
        v = weekly.get(lab, 0 if lab == "자사몰" else None)   # 자사몰 행만 있고 이번주 판매 없으면 0
        if v is None:
            continue
        cells.append((ri, v))
    print(f"  [매출데이터] {week} → 열 {col + 1}: " + ", ".join(f"{vals[ri][0]}={v:,}" for ri, v in cells))
    if dry:
        return

    if col + 1 > ws.col_count:
        ws.add_cols(col + 1 - ws.col_count)
    sid = ws.id
    prev = col - 1
    # 직전 주차 열 서식(테두리·#,##0) 복사 → 새 열
    reqs.append({"copyPaste": {
        "source": {"sheetId": sid, **gs_range(2, prev, last_row + 1, prev + 1)},
        "destination": {"sheetId": sid, **gs_range(2, col, last_row + 1, col + 1)},
        "pasteType": "PASTE_FORMAT"}})
    # 최신 열만 연노랑, 나머지 주차 열은 흰색
    reqs.append({"repeatCell": {"range": {"sheetId": sid, **gs_range(2, 1, last_row + 1, col)},
                                "cell": {"userEnteredFormat": {"backgroundColor": WHITE}},
                                "fields": "userEnteredFormat.backgroundColor"}})
    reqs.append({"repeatCell": {"range": {"sheetId": sid, **gs_range(2, col, last_row + 1, col + 1)},
                                "cell": {"userEnteredFormat": {"backgroundColor": LATEST}},
                                "fields": "userEnteredFormat.backgroundColor"}})
    sh.batch_update({"requests": reqs})
    reqs.clear()
    a1 = lambda r, c: openpyxl.utils.get_column_letter(c + 1) + str(r + 1)
    data = [{"range": a1(2, col), "values": [[week]]}] + \
           [{"range": a1(ri, col), "values": [[v]]} for ri, v in cells]
    ws.batch_update(data, value_input_option="RAW")


def col_widths(sh, tab):
    """탭 전체 열 너비(px) 목록. (범위를 탭 전체로 줘야 모든 열 메타가 옴)"""
    meta = sh.fetch_sheet_metadata({"ranges": [f"'{tab}'"], "fields": "sheets(data.columnMetadata.pixelSize)"})
    return [m.get("pixelSize", 100) for m in meta["sheets"][0]["data"][0].get("columnMetadata", [])]


def width_requests(sh, ws, tab, ref, start, width):
    """새 블록(start~)의 열 너비를 직전 블록(ref~)과 같게, 블록 앞뒤 빈 열은 기존 빈 열 너비로."""
    widths = col_widths(sh, tab)
    gap = widths[ref - 1] if ref > 0 else (widths[ref + width] if ref + width < len(widths) else 50)
    plan = {start + k: widths[ref + k] for k in range(width) if ref + k < len(widths)}
    if start > 0:
        plan[start - 1] = gap          # 블록 앞 빈 열
    plan[start + width] = gap          # 블록 뒤 빈 열
    return [{"updateDimensionProperties": {
        "range": {"sheetId": ws.id, "dimension": "COLUMNS", "startIndex": c, "endIndex": c + 1},
        "properties": {"pixelSize": w}, "fields": "pixelSize"}} for c, w in plan.items() if c < ws.col_count]


PRICE_HEADERS = {"금액", "결제금액", "총합계"}


def price_format_requests(ws, vals):
    """탭의 모든 블록에서 헤더(4행)가 금액류인 열의 데이터 행에 천 단위 쉼표(#,##0) 적용."""
    if len(vals) < 5:
        return []
    last = max((i for i, r in enumerate(vals) if any(str(x).strip() for x in r)), default=4) + 1
    return [{"repeatCell": {
        "range": {"sheetId": ws.id, **gs_range(4, c, last, c + 1)},
        "cell": {"userEnteredFormat": {"numberFormat": {"type": "NUMBER", "pattern": "#,##0"}}},
        "fields": "userEnteredFormat.numberFormat"}}
        for c, h in enumerate(vals[3]) if str(h).strip() in PRICE_HEADERS]


def option_align_requests(ws, vals):
    """모든 블록의 '옵션' 열 데이터 행을 왼쪽 정렬(숫자만 있는 옵션도 왼쪽)."""
    if len(vals) < 5:
        return []
    last = max((i for i, r in enumerate(vals) if any(str(x).strip() for x in r)), default=4) + 1
    return [{"repeatCell": {
        "range": {"sheetId": ws.id, **gs_range(4, c, last, c + 1)},
        "cell": {"userEnteredFormat": {"horizontalAlignment": "LEFT"}},
        "fields": "userEnteredFormat.horizontalAlignment"}}
        for c, h in enumerate(vals[3]) if str(h).strip() == "옵션"]


def tab_format_requests(ws, vals):
    return price_format_requests(ws, vals) + option_align_requests(ws, vals)


def data_format_request(ws, ref, start, width, nrows):
    """직전 블록 첫 데이터 행(5행)의 서식(정렬·글꼴·배경 등)을 새 블록 데이터 행 전체에 복사."""
    return {"copyPaste": {
        "source": {"sheetId": ws.id, **gs_range(4, ref, 5, ref + width)},
        "destination": {"sheetId": ws.id, **gs_range(4, start, 4 + nrows, start + width)},
        "pasteType": "PASTE_FORMAT"}}


def push_rank(sh, ws, wb, tab, week, overwrite, dry):
    vals = ws.get_all_values()
    row3 = vals[2] if len(vals) > 2 else []
    starts = [i for i, v in enumerate(row3) if v.strip()]
    if not starts:
        print(f"  [{tab}] 기존 블록이 없어 양식을 알 수 없음 → 건너뜀")
        return
    hit = [i for i in starts if week_key(row3[i]) == week_key(week)]
    if hit and not overwrite:
        print(f"  [{tab}] {week} 이미 있음 → 건너뜀")
        return
    ref = [s for s in starts if not hit or s != hit[0]][-1]      # 서식을 복사할 직전 블록
    stride = (starts[-1] - starts[-2]) if len(starts) >= 2 else None
    width = stride - 1 if stride else len([h for h in vals[3][ref:] if True])
    start = hit[0] if hit else starts[-1] + stride

    src = wb[tab] if tab in wb.sheetnames else None
    rows = []
    if src is not None:
        ncol = max(c for c in range(1, src.max_column + 1) if src.cell(3, c).value not in (None, ""))
        if ncol != width:
            print(f"  [{tab}] 열 개수 불일치(주간파일 {ncol} / 시트 블록 {width}) → 건너뜀")
            return
        codes = {c for c in range(1, ncol + 1) if str(src.cell(3, c).value or "").strip() in CODE_HEADERS}
        for r in range(4, src.max_row + 1):
            row = [clean(src.cell(r, c).value, c in codes) for c in range(1, ncol + 1)]
            if any(x != "" for x in row):
                rows.append(row)
    print(f"  [{tab}] {week} → 블록 시작 열 {openpyxl.utils.get_column_letter(start + 1)}, "
          f"폭 {width}, 데이터 {len(rows)}행" + ("" if src is not None else " (이번 주 순위 없음 → 빈 블록)"))
    if dry:
        return

    need = start + width + 1
    if need > ws.col_count:
        ws.add_cols(need - ws.col_count)
    if len(rows) + 5 > ws.row_count:
        ws.add_rows(len(rows) + 5 - ws.row_count)
    sid = ws.id
    reqs = []
    if hit:  # 덮어쓰기: 기존 블록 값 비우기
        reqs.append({"updateCells": {"range": {"sheetId": sid, **gs_range(4, start, ws.row_count, start + width)},
                                     "fields": "userEnteredValue"}})
    # 직전 블록 3~4행(라벨 서식·병합·헤더) 복사
    reqs.append({"copyPaste": {
        "source": {"sheetId": sid, **gs_range(2, ref, 4, ref + width)},
        "destination": {"sheetId": sid, **gs_range(2, start, 4, start + width)},
        "pasteType": "PASTE_NORMAL"}})
    if rows:
        reqs.append(data_format_request(ws, ref, start, width, len(rows)))
        # 직전 블록 데이터 행에 병합이 있으면 서식 복사 때 같이 따라오므로 데이터 영역 병합은 해제
        reqs.append({"unmergeCells": {"range": {"sheetId": sid, **gs_range(4, start, 4 + len(rows), start + width)}}})
    reqs += width_requests(sh, ws, tab, ref, start, width)
    sh.batch_update({"requests": reqs})

    a1 = openpyxl.utils.get_column_letter
    data = [{"range": f"{a1(start + 1)}3", "values": [[week]]}]
    if rows:
        data.append({"range": f"{a1(start + 1)}5:{a1(start + width)}{4 + len(rows)}", "values": rows})
    ws.batch_update(data, value_input_option="RAW")
    fmt = tab_format_requests(ws, ws.get_all_values())       # 금액 쉼표 + 옵션 왼쪽 정렬
    if fmt:
        sh.batch_update({"requests": fmt})


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--company", required=True, choices=list(SHEET_IDS))
    ap.add_argument("--file", help="주간판매데이터 xlsx")
    ap.add_argument("--format-only", action="store_true", help="입력 없이 모든 순위 탭 서식만 정리(금액 쉼표·옵션 왼쪽 정렬)")
    ap.add_argument("--week", help="주차 라벨(생략 시 파일명에서)")
    ap.add_argument("--overwrite", action="store_true", help="이미 있는 주차도 다시 쓰기")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    if a.format_only:
        sh = open_sheet("https://docs.google.com/spreadsheets/d/" + SHEET_IDS[a.company])
        for ws in sh.worksheets():
            if ws.title in RANK_TABS:
                vals = ws.get_all_values()
                pr, op = price_format_requests(ws, vals), option_align_requests(ws, vals)
                if pr or op:
                    sh.batch_update({"requests": pr + op})
                print(f"  [{ws.title}] 금액 열 {len(pr)}개 #,##0 / 옵션 열 {len(op)}개 왼쪽 정렬")
        return
    if not a.file:
        ap.error("--file 이 필요합니다")
    week = a.week or week_from_file(a.file)
    wb = openpyxl.load_workbook(a.file, data_only=True)
    sh = open_sheet("https://docs.google.com/spreadsheets/d/" + SHEET_IDS[a.company])
    print(f"[{a.company}] {sh.title} ← {os.path.basename(a.file)} ({week})" + (" [DRY-RUN]" if a.dry_run else ""))
    tabs = {w.title: w for w in sh.worksheets()}
    push_summary(sh, tabs["매출데이터"], wb, week, a.overwrite, a.dry_run, [])
    for tab in RANK_TABS:
        if tab in tabs:
            push_rank(sh, tabs[tab], wb, tab, week, a.overwrite, a.dry_run)
    print("완료" if not a.dry_run else "계획 출력만 함(쓰지 않음)")


if __name__ == "__main__":
    main()
