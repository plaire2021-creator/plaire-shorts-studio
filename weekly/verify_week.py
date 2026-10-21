# -*- coding: utf-8 -*-
"""
주간 판매데이터 검증 — 원본 ↔ 주간파일 ↔ 구글 시트 대조 (오류율 0% 확인용)
------------------------------------------------------------------
generate_weekly_sales.py 의 집계 함수를 쓰지 않고 원본을 직접 다시 계산해서 비교한다
(같은 코드로 검증하면 같은 버그를 못 잡으므로).

  [0] 기간    : 네이버 파일명 기간 / 오늘의집 주문결제완료일 / 자사몰 발주일이 같은 주인지,
                쿠팡 합계가 지난주 파일과 똑같지 않은지(지난주 파일 재사용 의심)
  [1] 합계    : 원본 채널 합계 ↔ 주간파일 매출데이터(이번 주 열)
  [2] 상세    : 원본 항목별 금액·수량 ↔ 주간파일 순위표
                네이버=채널상품번호별, 쿠팡=옵션ID별, 오늘의집=상품명별, 자사몰=합계·건수
                + 순위표의 총합계 = 그룹 옵션 금액 합
  [3] 시트    : 주간파일 ↔ 구글 시트(매출데이터 열 + 각 순위 블록 칸 단위)

실행:
  python verify_week.py --company 플레르가구 --file "2026_플레르가구 주간판매데이터 (10월 12주차).xlsx" \
      --naver <sales_*.xlsx> --coupang <SELLER_*.xlsx> --ohouse <주문배송*.xlsx> [--cafe24 <plaire1054_*.csv>]
  --no-sheet : 구글 시트 대조 생략 / --sheet-only : 원본 없이 주간파일↔시트만
종료코드 0 = 오류 0건, 1 = 불일치 있음.
"""
import argparse
import csv
import glob
import io
import os
import re
import sys
from collections import defaultdict
from datetime import date

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from generate_weekly_sales import DEFAULT_OHOUSE_PW as OHOUSE_PW  # noqa: E402  (각 PC의 .secrets 에서 읽음)


# ---------------------------------------------------------------- 공통
def n(v):
    try:
        return float(str(v).replace(",", "").strip())
    except (ValueError, TypeError):
        return 0.0


def norm(v):
    """비교용 문자열: 콤마 제거, 정수형 float → int, None → ''."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace(",", "").strip()


def hmap(ws, row=1):
    return {str(ws.cell(row, c).value).strip(): c for c in range(1, ws.max_column + 1)
            if ws.cell(row, c).value is not None}


class Report:
    def __init__(self):
        self.sections = []

    def section(self, title):
        self.sections.append({"title": title, "checked": 0, "errors": []})

    def check(self, ok, what):
        s = self.sections[-1]
        s["checked"] += 1
        if not ok:
            s["errors"].append(what)

    def print(self):
        tot = sum(s["checked"] for s in self.sections)
        err = sum(len(s["errors"]) for s in self.sections)
        print("\n================ 검증 결과 ================")
        for s in self.sections:
            mark = "OK " if not s["errors"] else "ERR"
            print(f"[{mark}] {s['title']}: {s['checked']}건 대조, 불일치 {len(s['errors'])}건")
            for e in s["errors"][:15]:
                print("       -", e)
            if len(s["errors"]) > 15:
                print(f"       ... 외 {len(s['errors']) - 15}건")
        rate = (err / tot * 100) if tot else 0.0
        print(f"--------------------------------------------\n총 {tot}건 대조 / 불일치 {err}건 / 오류율 {rate:.2f}%")
        return err


# ---------------------------------------------------------------- 원본 다시 계산 (독립)
def raw_naver(path):
    ws = openpyxl.load_workbook(path, data_only=True)
    ws = ws["SALES"] if "SALES" in ws.sheetnames else ws.worksheets[0]
    h = hmap(ws)
    per = defaultdict(lambda: {"금액": 0.0, "결제수": 0.0, "결제상품수량": 0.0})
    total = 0.0
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, h["채널상품명"]).value
        if nm is None or str(nm).strip() == "전체":
            continue
        code = str(ws.cell(r, h["채널상품번호"]).value).strip()
        amt = n(ws.cell(r, h["판매금액(순)"]).value)
        total += amt
        per[code]["금액"] += amt
        per[code]["결제수"] += n(ws.cell(r, h["상품결제건수"]).value)
        per[code]["결제상품수량"] += n(ws.cell(r, h["결제상품수량"]).value)
    return total, dict(per)


def raw_coupang(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["vendor item metrics"] if "vendor item metrics" in wb.sheetnames else wb.worksheets[0]
    h = hmap(ws)
    oid_c = h.get("옵션 ID") or h.get("옵션ID")
    per, total = {}, 0.0
    for r in range(2, ws.max_row + 1):
        oid = ws.cell(r, oid_c).value
        if oid is None:
            continue
        amt = n(ws.cell(r, h["매출(원)"]).value)
        total += amt
        if amt > 0:
            per[str(oid).strip()] = {"결제금액": amt, "판매수": n(ws.cell(r, h["판매량"]).value)}
    return total, per


def raw_ohouse(path):
    import msoffcrypto
    buf = io.BytesIO()
    with open(path, "rb") as f:
        of = msoffcrypto.OfficeFile(f)
        if of.is_encrypted():
            of.load_key(password=OHOUSE_PW)
            of.decrypt(buf)
        else:
            buf.write(f.read())
    buf.seek(0)
    wb = openpyxl.load_workbook(buf, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    h = hmap(ws)
    amt_c = h.get("판매가 * 수량") or h.get("판매가*수량")
    per = defaultdict(lambda: {"결제금액": 0.0, "건수": 0, "ids": set()})
    total, dates = 0.0, []
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, h["상품명"]).value
        if nm is None:
            continue
        amt = n(ws.cell(r, amt_c).value)
        total += amt
        key = re.sub(r"^\s*\d+\s*%\s*즉시할인\s*[ㅣ|/]?\s*", "", str(nm).strip()).strip()
        per[key]["결제금액"] += amt
        per[key]["건수"] += 1
        if "상품아이디" in h:
            per[key]["ids"].add(norm(ws.cell(r, h["상품아이디"]).value))
        d = ws.cell(r, h["주문결제완료일"]).value if "주문결제완료일" in h else None
        if d:
            dates.append(str(d)[:10])
    return total, dict(per), dates


def raw_cafe24(path):
    """자사몰 확정 규칙(cafe24.py와 동일 정의): 주문 단위 F열 총 주문금액, G열 총 결제금액 0원 주문은 취소로 제외."""
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if rows and "판매금액" in rows[0]:                      # '상품별 매출 분석' 형식
        return sum(n(r["판매금액"]) for r in rows), sum(1 for r in rows if n(r["판매금액"]) > 0), []
    orders = defaultdict(lambda: {"F": 0.0, "G": 0.0, "lines": 0})
    dates = []
    for r in rows:
        o = orders[r.get("주문번호", "")]
        o["F"] = max(o["F"], n(r.get("총 주문금액")))
        o["G"] = max(o["G"], n(r.get("총 결제금액")))
        o["lines"] += 1
        if r.get("발주일"):
            dates.append(r["발주일"][:10])
    live = [o for o in orders.values() if o["G"] > 0]
    return sum(o["F"] for o in live), sum(o["lines"] for o in live), dates


# ---------------------------------------------------------------- 주간파일 읽기
def weekly_summary(wb):
    s = wb["매출데이터"]
    return {str(s.cell(r, 1).value).strip(): n(s.cell(r, 2).value)
            for r in range(4, 12) if s.cell(r, 1).value}


def weekly_rank(wb, tab):
    """순위 시트 → (헤더, 행 목록[dict], 그룹 목록)."""
    if tab not in wb.sheetnames:
        return None, [], []
    ws = wb[tab]
    ncol = max(c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value not in (None, ""))
    hdr = [str(ws.cell(3, c).value or "").strip() for c in range(1, ncol + 1)]
    rows, groups = [], []
    for r in range(4, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ncol + 1)]
        if all(v in (None, "") for v in vals):
            continue
        d = dict(zip(hdr, vals))
        d["_raw"] = vals
        rows.append(d)
        if vals[1] not in (None, ""):
            groups.append({"name": vals[1], "total": d.get("총합계"), "rows": []})
        if groups:
            groups[-1]["rows"].append(d)
    return hdr, rows, groups


# ---------------------------------------------------------------- 검증
def week_range_from_naver(path):
    m = re.search(r"(\d{4})(\d{2})(\d{2})-(\d{4})(\d{2})(\d{2})", os.path.basename(path))
    if not m:
        return None
    y1, m1, d1, y2, m2, d2 = map(int, m.groups())
    return date(y1, m1, d1).isoformat(), date(y2, m2, d2).isoformat()


def prev_weekly(company, path):
    def key(p):
        m = re.search(r"\((\d+)\s*월\s*(\d+)\s*주차\)", p)
        return (int(m.group(1)), int(m.group(2))) if m else (0, 0)
    me = key(os.path.basename(path))
    cands = [p for p in glob.glob(os.path.join(HERE, f"2026_{company} 주간판매데이터 (*주차).xlsx"))
             if not os.path.basename(p).startswith("~$") and key(os.path.basename(p)) < me]
    return max(cands, key=lambda p: key(os.path.basename(p))) if cands else None


def run(a):
    rep = Report()
    wb = openpyxl.load_workbook(a.file, data_only=True)
    summ = weekly_summary(wb)

    if a.sheet_only:
        return sheet_check(rep, a, wb, summ)
    nv_total, nv_per = raw_naver(a.naver)
    cp_total, cp_per = raw_coupang(a.coupang)
    oh_total, oh_per, oh_dates = raw_ohouse(a.ohouse)
    c24 = raw_cafe24(a.cafe24) if a.cafe24 else None

    # [0] 기간
    rep.section("[0] 파일 기간")
    rng = week_range_from_naver(a.naver)
    rep.check(rng is not None, "네이버 파일명에서 기간을 못 읽음")
    if rng:
        lo, hi = rng
        bad = [d for d in oh_dates if not (lo <= d <= hi)]
        rep.check(not bad, f"오늘의집 주문일이 주간({lo}~{hi}) 밖: {sorted(set(bad))[:5]}")
        if c24 and c24[2]:
            bad = [d for d in c24[2] if not (lo <= d <= hi)]
            rep.check(not bad, f"자사몰 발주일이 주간({lo}~{hi}) 밖: {sorted(set(bad))[:5]}")
    pw = prev_weekly(a.company, a.file)
    if pw:
        prev = weekly_summary(openpyxl.load_workbook(pw, data_only=True))
        rep.check(abs(prev.get("쿠팡", -1) - cp_total) >= 1,
                  f"쿠팡 합계 {cp_total:,.0f}가 지난주({os.path.basename(pw)})와 같음 → 지난주 파일 재사용 의심")

    # [1] 합계
    rep.section("[1] 원본 합계 ↔ 주간파일 매출데이터")
    exp = {"네이버": nv_total, "쿠팡": cp_total, "오늘의집": oh_total}
    if c24:
        exp["자사몰"] = c24[0]
    for ch, v in exp.items():
        rep.check(abs(summ.get(ch, 0) - v) < 1, f"{ch}: 원본 {v:,.0f} / 주간파일 {summ.get(ch, 0):,.0f}")
    if "자사몰" in summ and not c24:
        rep.check(summ["자사몰"] == 0, f"자사몰 원본 없음인데 주간파일 {summ['자사몰']:,.0f}")
    rep.check(abs(summ.get("합계", 0) - sum(exp.values())) < 1,
              f"합계: 원본 {sum(exp.values()):,.0f} / 주간파일 {summ.get('합계', 0):,.0f}")

    # [2] 상세
    rep.section("[2] 원본 항목별 ↔ 주간파일 순위표")
    # 네이버: 채널상품번호별
    _, rows, groups = weekly_rank(wb, "네이버")
    got = {norm(r["상품 코드"]): r for r in rows}
    rep.check(set(got) == set(nv_per), f"네이버 상품코드 목록 차이: 원본만 {sorted(set(nv_per) - set(got))[:5]}, "
                                       f"주간파일만 {sorted(set(got) - set(nv_per))[:5]}")
    for code, v in nv_per.items():
        r = got.get(code)
        if r is None:
            continue
        for k in ("금액", "결제수", "결제상품수량"):
            rep.check(abs(n(r[k]) - v[k]) < 1, f"네이버 {code} {k}: 원본 {v[k]:,.0f} / 주간파일 {n(r[k]):,.0f}")
    # 쿠팡: 옵션ID별
    _, rows, g2 = weekly_rank(wb, "쿠팡")
    groups += g2
    got = {norm(r["옵션코드"]): r for r in rows}
    rep.check(set(got) == set(cp_per), f"쿠팡 옵션ID 목록 차이: 원본만 {sorted(set(cp_per) - set(got))[:5]}, "
                                       f"주간파일만 {sorted(set(got) - set(cp_per))[:5]}")
    for oid, v in cp_per.items():
        r = got.get(oid)
        if r is None:
            continue
        for k in ("결제금액", "판매수"):
            rep.check(abs(n(r[k]) - v[k]) < 1, f"쿠팡 {oid} {k}: 원본 {v[k]:,.0f} / 주간파일 {n(r[k]):,.0f}")
    # 오늘의집: 상품명별
    _, rows, g3 = weekly_rank(wb, "오늘의집")
    groups += g3
    by_name = {str(g["name"]).strip(): g for g in g3}
    # 식탁·테이블·책상은 사이즈 무관 한 제품(제품명 뒤 사이즈 제거) → 원본 상품명을 주간파일 제품명(앞부분 일치)으로 합침
    merged = defaultdict(lambda: {"결제금액": 0.0, "건수": 0, "ids": set()})
    unmatched = []
    for nm, v in oh_per.items():
        key = nm if nm in by_name else max((w for w in by_name if nm.startswith(w)), key=len, default=None)
        if key is None:
            unmatched.append(nm)
            continue
        m = merged[key]
        m["결제금액"] += v["결제금액"]; m["건수"] += v["건수"]; m["ids"] |= v["ids"]
    rep.check(not unmatched and set(merged) == set(by_name),
              f"오늘의집 상품명 목록 차이: 원본만 {unmatched[:3]}, 주간파일만 {sorted(set(by_name) - set(merged))[:3]}")
    for nm, v in merged.items():
        g = by_name.get(nm)
        if g is None:
            continue
        rep.check(abs(n(g["total"]) - v["결제금액"]) < 1, f"오늘의집 '{nm}' 금액: 원본 {v['결제금액']:,.0f} / 주간파일 {n(g['total']):,.0f}")
        cnt = sum(n(r.get("결제수")) for r in g["rows"])
        rep.check(cnt == v["건수"], f"오늘의집 '{nm}' 결제수: 원본 {v['건수']} / 주간파일 {cnt:.0f}")
        gid = norm(g["rows"][0].get("상품아이디"))
        rep.check(not v["ids"] or gid in v["ids"], f"오늘의집 '{nm}' 상품아이디 {gid} 가 원본({sorted(v['ids'])})에 없음")
    # 자사몰: 합계·건수
    _, rows, g4 = weekly_rank(wb, "자사몰")
    groups += g4
    if c24:
        rep.check(abs(sum(n(r["금액"]) for r in rows) - c24[0]) < 1,
                  f"자사몰 순위표 금액 합 {sum(n(r['금액']) for r in rows):,.0f} / 원본 {c24[0]:,.0f}")
        rep.check(sum(n(r["결제수"]) for r in rows) == c24[1],
                  f"자사몰 결제수 합 {sum(n(r['결제수']) for r in rows):.0f} / 원본 주문라인 {c24[1]}")
    else:
        rep.check(not rows, "자사몰 원본이 없는데 주간파일에 자사몰 순위가 있음")
    # 그룹 총합계 = 옵션 금액 합 (모든 채널)
    for g in groups:
        amt_key = "금액" if "금액" in g["rows"][0] else "결제금액"
        s = sum(n(r.get(amt_key)) for r in g["rows"])
        rep.check(abs(n(g["total"]) - s) < 1, f"'{g['name']}' 총합계 {n(g['total']):,.0f} ≠ 옵션 합 {s:,.0f}")

    if not a.no_sheet:
        return sheet_check(rep, a, wb, summ)
    return rep.print()


def sheet_check(rep, a, wb, summ):
    if True:
        rep.section("[3] 주간파일 ↔ 구글 시트")
        sys.path.insert(0, os.path.dirname(HERE))
        from sheets_client import open_sheet
        import push_gsheet as pg
        week = pg.week_from_file(a.file)
        sh = open_sheet("https://docs.google.com/spreadsheets/d/" + pg.SHEET_IDS[a.company])
        tabs = {w.title: w for w in sh.worksheets()}
        sv = tabs["매출데이터"].get_all_values()
        cols = [i for i, v in enumerate(sv[2]) if i > 0 and pg.week_key(v) == pg.week_key(week)]
        rep.check(len(cols) == 1, f"매출데이터에 {week} 열이 {len(cols)}개")
        if cols:
            c = cols[0]
            for r in sv[3:]:
                lab = r[0].strip() if r else ""
                if lab in summ:
                    rep.check(norm(r[c]) == norm(summ[lab]), f"매출데이터 {lab}: 시트 {r[c]} / 주간파일 {summ[lab]:,.0f}")
        for tab in pg.RANK_TABS:
            if tab not in tabs:
                continue
            hdr, rows, _ = weekly_rank(wb, tab)
            vals = tabs[tab].get_all_values()
            starts = [i for i, v in enumerate(vals[2]) if pg.week_key(v) == pg.week_key(week)]
            rep.check(len(starts) == 1, f"{tab}: {week} 블록이 {len(starts)}개")
            if len(starts) != 1:
                continue
            s0 = starts[0]
            w = len(hdr) if hdr else 0
            if not w:                                  # 이번 주 순위 없음 → 시트 블록도 비어 있어야
                filled = [i for i, r in enumerate(vals[4:], 5) if any(x.strip() for x in r[s0:s0 + 6])]
                rep.check(not filled, f"{tab}: 이번 주 순위가 없는데 시트 블록 {filled[:3]}행에 값이 있음")
                continue
            for i, r in enumerate(rows):
                want = [norm(pg.clean(v)) for v in r["_raw"]]
                got = [norm(x) for x in (vals[4 + i][s0:s0 + w] if 4 + i < len(vals) else [])]
                got += [""] * (w - len(got))
                rep.check(want == got, f"{tab} {5 + i}행: 주간파일 {want} / 시트 {got}")
            extra = [5 + i for i, r in enumerate(vals[4 + len(rows):], len(rows))
                     if any(x.strip() for x in r[s0:s0 + w])]
            rep.check(not extra, f"{tab}: 블록 아래에 남은 값 있음(행 {extra[:5]})")
    return rep.print()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--company", required=True, choices=["플레르가구", "가즈다가구"])
    ap.add_argument("--file", required=True)
    ap.add_argument("--naver")
    ap.add_argument("--coupang")
    ap.add_argument("--ohouse")
    ap.add_argument("--cafe24")
    ap.add_argument("--no-sheet", action="store_true")
    ap.add_argument("--sheet-only", action="store_true", help="원본 없이 주간파일↔구글 시트만 대조")
    a = ap.parse_args()
    if not a.sheet_only and not (a.naver and a.coupang and a.ohouse):
        ap.error("--naver --coupang --ohouse 가 필요합니다 (원본이 없으면 --sheet-only)")
    sys.exit(1 if run(a) else 0)


if __name__ == "__main__":
    main()
