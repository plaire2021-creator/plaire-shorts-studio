"""신제품 출시 매뉴얼 '성과': 주간 매출 통계 결과 파일에서 기간 안의 주차별·채널별 판매를 상품 단위로 꺼낸다.

  python launch_sales.py <주간 매출 통계 폴더> <브랜드> <시작 YYYY-MM-DD> <끝 YYYY-MM-DD>   → 마지막 줄에 JSON

파일: '2026_플레르가구 주간판매데이터 (10월 12주차).xlsx' — 시트 네이버·쿠팡·오늘의집, 각 시트 A1에 '2026.10.12~2026.10.18'
상품 한 덩어리 = 순위가 적힌 줄부터 다음 순위 전까지(옵션 줄 포함). 덩어리마다 그 안의 상품코드(ID)들과 수량·금액 합계를 낸다.
  네이버: 제품명 | 옵션 | 상품 코드 | 결제수 | 결제상품수량 | 금액
  쿠팡:   상품명 | 옵션 | 옵션코드(vendorItemId) | 결제금액 | 판매수
  오늘의집: 상품명 | 상품아이디 | 옵션 | 결제수 | 결제금액
"""
import glob
import json
import os
import re
import sys

import openpyxl

# 시트 이름 → (이름 열, ID 열, 수량 열, 금액 열) — 0부터, 순위 열(0) 다음부터
COLS = {"네이버": (1, 3, 5, 6), "쿠팡": (1, 3, 5, 4), "오늘의집": (1, 2, 4, 5)}


def num(v):
    try:
        return float(str(v).replace(",", "")) if v not in (None, "") else 0
    except ValueError:
        return 0


def period(ws):
    m = re.search(r"(\d{4})\.(\d{1,2})\.(\d{1,2})\s*~\s*(\d{4})\.(\d{1,2})\.(\d{1,2})", str(ws.cell(1, 1).value or ""))
    if not m:
        return None
    y1, m1, d1, y2, m2, d2 = map(int, m.groups())
    return f"{y1:04d}-{m1:02d}-{d1:02d}", f"{y2:04d}-{m2:02d}-{d2:02d}"


def blocks(ws, cols):
    ni, ii, qi, ai = cols
    out, cur = [], None
    for r in ws.iter_rows(min_row=4, values_only=True):
        r = list(r) + [None] * 10
        rank, name = r[0], r[ni]
        if rank not in (None, "") and str(rank).strip().isdigit():
            cur = {"name": str(name or "").strip(), "ids": [], "qty": 0, "amount": 0}
            out.append(cur)
        if cur is None:
            continue
        if r[ii] not in (None, ""):
            cur["ids"].append(re.sub(r"\.0$", "", str(r[ii]).strip()))
        cur["qty"] += num(r[qi])
        cur["amount"] += num(r[ai])
    return out


def main():
    folder, brand, start, end = sys.argv[1:5]
    weeks = []
    for f in glob.glob(os.path.join(folder, f"*{brand}*주간판매데이터*.xlsx")):
        if os.path.basename(f).startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        except Exception:
            continue
        p = period(wb["네이버"]) if "네이버" in wb.sheetnames else None
        if not p or p[1] < start or p[0] > end:
            continue
        w = {"start": p[0], "end": p[1], "file": os.path.basename(f), "ch": {}}
        for sh, cols in COLS.items():
            if sh in wb.sheetnames:
                w["ch"][sh] = blocks(wb[sh], cols)
        weeks.append(w)
    weeks.sort(key=lambda w: w["start"])
    print(json.dumps({"weeks": weeks}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(json.dumps({"error": f"주간 매출 파일을 읽지 못했어요: {e}"}, ensure_ascii=False))
