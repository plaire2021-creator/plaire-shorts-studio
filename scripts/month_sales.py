"""홈 '월간 목표 매출' 카드: 주간 매출 통계 결과 파일로 이번 달 브랜드·채널별 매출을 모은다.

  python month_sales.py <주간 매출 통계 폴더> <YYYY-MM>   → 마지막 줄에 JSON

파일마다 그 주 합계 = '매출데이터' 시트 첫 주 열(B열), 기간 = '네이버' 시트 A1.
한 주가 두 달에 걸치면 그 달에 든 날짜 비율만큼 나눠 넣는다(주 단위 자료라 날짜별 매출은 없음).
채널: 네이버(=스마트스토어)·쿠팡·오늘의집 (월간 목표가 이 세 채널 기준이라 자사몰은 따로 둔다)
"""
import calendar
import datetime as dt
import glob
import json
import os
import re
import sys

import openpyxl

CH = {"네이버": "스마트스토어", "쿠팡": "쿠팡", "오늘의집": "오늘의집", "자사몰": "자사몰"}


def num(v):
    try:
        return float(str(v).replace(",", "")) if v not in (None, "") else 0.0
    except ValueError:
        return 0.0


def main():
    folder, month = sys.argv[1], sys.argv[2]
    y, m = map(int, month.split("-"))
    m0, m1 = dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])
    out = {}
    last = {}
    for f in glob.glob(os.path.join(folder, "*주간판매데이터*.xlsx")):
        base = os.path.basename(f)
        if base.startswith("~$"):
            continue
        brand = "플레르가구" if "플레르" in base else "가즈다가구" if "가즈다" in base else None
        if not brand:
            continue
        try:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            p = re.search(r"(\d{4})\.(\d{1,2})\.(\d{1,2})\s*~\s*(\d{4})\.(\d{1,2})\.(\d{1,2})", str(wb["네이버"].cell(1, 1).value or ""))
            if not p:
                continue
            a = dt.date(*map(int, p.groups()[:3]))
            b = dt.date(*map(int, p.groups()[3:]))
            days = (b - a).days + 1
            inside = (min(b, m1) - max(a, m0)).days + 1
            if inside <= 0:
                continue
            share = inside / days
            rows = list(wb["매출데이터"].iter_rows(min_row=1, max_row=15, values_only=True))
        except Exception:
            continue
        o = out.setdefault(brand, {v: 0.0 for v in CH.values()})
        for r in rows:
            if r and r[0] in CH and len(r) > 1:
                o[CH[r[0]]] += num(r[1]) * share
        last[brand] = max(last.get(brand, ""), b.isoformat())
    print(json.dumps({"month": month, "actual": {k: {c: round(v) for c, v in d.items()} for k, d in out.items()}, "through": last}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(json.dumps({"error": f"주간 매출 파일을 읽지 못했어요: {e}"}, ensure_ascii=False))
