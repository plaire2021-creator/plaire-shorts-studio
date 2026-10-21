"""스마트스토어 리뷰 엑셀에서 5점이 아닌 리뷰만 남긴 '리뷰 수정요청' 엑셀을 만든다.
원본과 같은 열·머리글 서식을 그대로 두고, 평점 5점 행과 빈 행만 지운다.

  python low_reviews.py <리뷰엑셀> <브랜드: 플레르가구|가즈다가구> <저장폴더>
출력(마지막 줄 JSON): {"file": 저장경로, "total": 전체 리뷰 수, "low": 5점 미만 수}
"""
import json
import os
import re
import sys
from datetime import date

import openpyxl


def main():
    src, brand, out_dir = sys.argv[1], sys.argv[2], sys.argv[3]
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    head = [str(c.value or "").strip() for c in ws[1]]
    if "구매자평점" not in head:
        sys.exit("리뷰 엑셀이 아니에요('구매자평점' 열이 없음). 스마트스토어 리뷰관리에서 받은 엑셀을 넣어 주세요.")
    col = head.index("구매자평점") + 1
    total = low = 0
    for r in range(ws.max_row, 1, -1):
        v = ws.cell(r, col).value
        empty = all(c.value in (None, "") for c in ws[r])
        if empty:
            ws.delete_rows(r)
            continue
        total += 1
        try:
            five = float(v) >= 5
        except (TypeError, ValueError):
            five = False
        if five:
            ws.delete_rows(r)
        else:
            low += 1
    m = re.search(r"review_(\d{8})", os.path.basename(src))
    day = m.group(1) if m else date.today().strftime("%Y%m%d")
    short = "플레르" if brand == "플레르가구" else "가즈다"
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, f"{short} 리뷰 수정요청_{day}.xlsx")
    wb.save(out)
    print(json.dumps({"file": out, "total": total, "low": low}, ensure_ascii=False))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
