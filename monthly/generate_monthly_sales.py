# -*- coding: utf-8 -*-
"""
월간 판매 데이터 자동 생성기
------------------------------------------------------------------
네이버 스마트스토어 / 쿠팡 / 오늘의집에서 **월 단위 기간**으로 받은 엑셀 3개를 넣으면
'월간판매데이터' 파일(매출데이터 + 통합순위 + 네이버/쿠팡/오늘의집 순위)을 만든다.

집계·그룹핑 규칙은 주간 시스템과 100% 동일하다. 파싱·순위표 로직을 복사하지 않고
`../주간 매출 통계/generate_weekly_sales.py` 를 그대로 불러 쓴다(규칙 이중관리 방지).

사용법:
    python generate_monthly_sales.py --input "플레르" --brand "플레르가구" \
        --base "2026_플레르가구 월간판매데이터 (7월).xlsx" \
        --out  "2026_플레르가구 월간판매데이터 (8월).xlsx"

주간과 다른 점:
  - 주차 라벨 대신 **월 라벨**('8월'), 매출데이터는 **1월→12월 왼→오 시간순 누적**(최대 12개월)
  - **통합순위 탭**: 3개 채널을 제품 기준으로 합친 월간 전체 판매 순위
"""
import argparse
import importlib.util
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

import merge_products as mp


# ----------------------------------------------------------------------
# 주간 스크립트를 '엔진'으로 로드 — 파싱·그룹핑·순위표 규칙의 단일 출처
# ----------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
ENGINE_CANDIDATES = [
    os.path.join(HERE, "generate_weekly_sales.py"),                      # 같은 폴더에 복사해 둔 경우
    os.path.join(os.path.dirname(HERE), "주간 매출 통계", "generate_weekly_sales.py"),
    os.path.join(HERE, "..", "주간 매출 통계", "generate_weekly_sales.py"),
]


def load_engine(explicit=None):
    paths = ([explicit] if explicit else []) + ENGINE_CANDIDATES
    for p in paths:
        if p and os.path.exists(p):
            spec = importlib.util.spec_from_file_location("weekly_engine", p)
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)
            return mod
    raise SystemExit(
        "generate_weekly_sales.py 를 찾지 못했습니다.\n"
        "  '주간 매출 통계' 폴더가 '월간 매출 통계'와 같은 상위 폴더에 있어야 합니다.\n"
        "  (또는 --engine \"경로\\generate_weekly_sales.py\" 로 직접 지정)")


# ----------------------------------------------------------------------
# 서식
# ----------------------------------------------------------------------
BLACK = Side(style="thin", color="000000")
BOX = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
THIN = Side(style="thin", color="BFBFBF")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
LATEST_FILL = PatternFill("solid", fgColor="FFF2CC")     # 최신 월 강조(연한 노랑)
NO_FILL = PatternFill(fill_type=None)
CAT_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(bold=True, size=14)
HDR_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

MAX_MONTHS = 12                                          # 매출데이터 누적 개월
CHANNELS = ["네이버", "쿠팡", "오늘의집", "자사몰"]       # 매출데이터 행 순서(주간과 동일)
# 통합순위에서 빼는 채널. 자사몰은 상품명이 사내 표기(견적건·결제창 등)라
# 타 채널 제품과 같은 기준으로 묶이지 않아 순위를 흐린다. 매출데이터·자사몰 탭에는 그대로 들어간다.
UNIFIED_EXCLUDE = {"자사몰"}


def num(v):
    try:
        return float(str(v).replace(",", "").strip())
    except (ValueError, TypeError, AttributeError):
        return 0.0


def month_key(label):
    m = re.search(r"(\d+)\s*월", str(label))
    return int(m.group(1)) if m else 99


def ohouse_quantities(gw, path, password=None):
    """오늘의집 제품별 결제수량 {제품명: 수량}.

    주간의 parse_ohouse 는 '결제수'(주문 줄 수)만 세고 수량은 안 담는다.
    통합순위의 결제수량을 채우려면 수량이 필요한데, 주간 엔진을 고치면 주간 결과에
    영향이 갈 수 있으므로 **엔진의 그룹핑 헬퍼를 그대로 빌려** 여기서 수량만 따로 센다.
    (제품명 키 산출은 parse_ohouse 와 동일한 규칙 — 결과는 아래에서 교차검증한다)
    """
    wb = gw.load_workbook_any(path, password=password)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    m = gw.header_map(ws)
    c_nm, c_qty = gw.col(m, "상품명"), m.get("수량")
    out = {}
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, c_nm).value
        if nm is None:
            continue
        base = gw.clean_ohouse_name(nm)
        key = gw._split_trailing_size(base)[0] if gw._is_table(base) else base
        qty = gw.num(ws.cell(r, c_qty).value) if c_qty else 1
        out[key] = out.get(key, 0) + (qty or 1)
    return out


def naver_quantities(gw, path):
    """네이버 제품별 **순수량** {제품명: 결제상품수량 − 환불상품수량}.

    네이버는 금액만 환불을 반영(판매금액 순)하고 수량은 결제 기준으로 준다.
    그대로 더하면 전액 환불된 제품이 '매출 0원 + 수량 1개'로 남는다.
    금액과 기준을 맞추기 위해 환불상품수량을 빼서 순수량으로 집계한다.
    제품 키는 parse_naver 와 같은 규칙(group_and_option)을 쓴다.
    """
    ws = gw.load_workbook_any(path)["SALES"]
    m = gw.header_map(ws)
    c_nm = gw.col(m, "채널상품명")
    c_q, c_rq = gw.col(m, "결제상품수량"), m.get("환불상품수량")
    out = {}
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, c_nm).value
        if nm is None or str(nm).strip() == "전체":      # 일별 '전체' 요약행 제외
            continue
        key = gw.group_and_option(str(nm).strip())[0]
        qty = gw.num(ws.cell(r, c_q).value)
        if c_rq:
            qty -= gw.num(ws.cell(r, c_rq).value)
        out[key] = out.get(key, 0) + qty
    return out


def attach_quantities(gw, naver_products, coupang_products, ohouse_products,
                      naver_path, ohouse_path, ohouse_pw):
    """각 채널 products에 '총수량'을 달아 준다. 채널마다 수량 열 이름이 다르다."""
    nmap = naver_quantities(gw, naver_path)        # 네이버: 결제 − 환불 (순수량)
    nmiss = [p["제품명"] for p in naver_products if p["제품명"] not in nmap]
    if nmiss:
        print(f"  [경고] 네이버 수량 매칭 실패 {len(nmiss)}건: {nmiss[:3]}")
    for p in naver_products:
        p["총수량"] = nmap.get(p["제품명"], 0)
    for p in coupang_products:                     # 쿠팡: 판매수
        p["총수량"] = sum(o.get("판매수", 0) for o in p["옵션들"])
    qmap = ohouse_quantities(gw, ohouse_path, ohouse_pw)
    missing = [p["제품명"] for p in ohouse_products if p["제품명"] not in qmap]
    if missing:                                    # 그룹핑 규칙이 어긋나면 바로 드러나게
        print(f"  [경고] 오늘의집 수량 매칭 실패 {len(missing)}건: {missing[:3]}")
    for p in ohouse_products:                      # 오늘의집: 원본 '수량' 열
        p["총수량"] = qmap.get(p["제품명"], 0)


def month_from_naver_name(path):
    """sales_20260801-20260831_... → '8월' (시작일의 월)."""
    m = re.search(r"(\d{4})(\d{2})(\d{2})-(\d{4})(\d{2})(\d{2})", os.path.basename(path))
    return f"{int(m.group(2))}월" if m else None


# ----------------------------------------------------------------------
# 매출데이터 탭 — 채널(행) × 월(열), 왼쪽=오래된 달
# ----------------------------------------------------------------------
def update_summary(ws, month_label, totals, brand):
    """기존 월 데이터를 모두 읽어 새 달을 얹고, 시간순으로 다시 그린다.
    (열 삽입 대신 전체 재작성 — 수식 밀림·라벨 뒤섞임이 원천적으로 없다)

    totals: {채널명: 금액}. 자사몰 행은 자사몰 매출이 있거나 base에 이미 있을 때만 넣는다
    (가즈다처럼 자사몰이 없는 회사는 3채널 그대로). 주간 양식과 동일한 배치.
    """
    # 기존 파일에서 채널이 어느 행에 있었는지 읽는다(구 3채널 파일도 그대로 인식)
    old_rows = {}
    for r in range(4, 9):
        lab = str(ws.cell(r, 1).value or "").strip()
        if lab in CHANNELS:
            old_rows[lab] = r

    # 값이 빈 칸이면 None으로 남긴다. '0원'과 '그 달은 데이터를 안 넣었음'은 다르다
    # (8월 월간을 자사몰 없이 만들어 놓고 9월부터 자사몰을 붙이는 경우가 실제로 있었다).
    def cellval(r, c):
        v = ws.cell(r, c).value
        return None if v in (None, "") else num(v)

    data = {}
    for c in range(2, ws.max_column + 1):
        lab = str(ws.cell(3, c).value or "").strip()
        if lab:
            data[lab] = {ch: cellval(r, c) for ch, r in old_rows.items()}
    data[month_label] = dict(totals)                      # 같은 달 재실행이면 덮어쓰기

    use_cafe24 = (totals.get("자사몰") or 0) > 0 or "자사몰" in old_rows
    chans = [c for c in CHANNELS if c != "자사몰" or use_cafe24]

    labels = sorted(data, key=month_key)[-MAX_MONTHS:]    # 오래된 달 왼쪽, 최대 12개월

    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 12),
                            min_col=1, max_col=max(ws.max_column, len(labels) + 2)):
        for cell in row:
            cell.value = None
            cell.fill = NO_FILL
            cell.border = Border()

    ws["A1"] = f"{brand} 월간판매 데이터"
    ws["A1"].font = TITLE_FONT
    sum_row = 4 + len(chans)
    for i, ch in enumerate(chans):
        ws.cell(4 + i, 1, ch)
    ws.cell(sum_row, 1, "합계")

    for i, lab in enumerate(labels):
        c = 2 + i
        vals = data[lab]
        ws.cell(3, c, lab).font = HDR_FONT
        ws.cell(3, c).alignment = CENTER
        for j, ch in enumerate(chans):
            v = vals.get(ch)
            cell = ws.cell(4 + j, c, None if v is None else int(round(v)))
            cell.number_format = "#,##0"
            cell.alignment = RIGHT
        cell = ws.cell(sum_row, c,
                       int(round(sum(vals.get(ch) or 0 for ch in chans))))
        cell.number_format = "#,##0"
        cell.alignment = RIGHT

    last = 1 + len(labels)
    for c in range(1, last + 1):
        L = openpyxl.utils.get_column_letter(c)
        ws.column_dimensions[L].width = 12 if c == 1 else 14
        ws.column_dimensions[L].hidden = False
        for r in range(3, sum_row + 1):
            cell = ws.cell(r, c)
            cell.border = BOX
            if c >= 2:
                cell.fill = LATEST_FILL if c == last else (HDR_FILL if r == 3 else NO_FILL)
            elif r == 3:
                cell.fill = HDR_FILL
    return labels, chans


# ----------------------------------------------------------------------
# 통합순위 탭 — 3개 채널 합산
# ----------------------------------------------------------------------
def write_unified_sheet(ws, period, rows, chans):
    headers = ["", "통합 제품명", "카테고리"] + list(chans) + ["합계", "결제수량"]
    n_ch = len(chans)
    c_sum = 4 + n_ch                                     # 합계 열
    c_qty = c_sum + 1                                    # 결제수량 열
    ws["A1"] = period
    ws["A1"].font = HDR_FONT
    ws["A2"] = "채널 합산 통합 판매순위"
    ws["A2"].font = TITLE_FONT
    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c, h)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = GRID

    for i, row in enumerate(rows):
        r = 4 + i
        ws.cell(r, 1, i + 1).alignment = CENTER
        ws.cell(r, 2, row["통합제품명"])
        ws.cell(r, 3, row["카테고리"]).alignment = CENTER
        ws.cell(r, 3).fill = CAT_FILL
        for c, key in [(4 + i, ch) for i, ch in enumerate(chans)] + [(c_sum, "합계")]:
            cell = ws.cell(r, c, int(round(row[key])))
            cell.number_format = "#,##0"
            cell.alignment = RIGHT
        ws.cell(r, c_sum).font = HDR_FONT
        qcell = ws.cell(r, c_qty, int(round(row["결제수량"])))
        qcell.number_format = "#,##0"
        qcell.alignment = CENTER
        for c in range(1, len(headers) + 1):
            ws.cell(r, c).border = GRID

    widths = {1: 6, 2: 55, 3: 14, c_sum: 14, c_qty: 10}
    for i in range(n_ch):
        widths.setdefault(4 + i, 13)
    for c, w in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "A4"


# ----------------------------------------------------------------------
# 빌드
# ----------------------------------------------------------------------
def build(gw, naver_path, coupang_path, ohouse_path, out_path, month_label,
          ohouse_pw=None, base_path=None, brand="플레르가구", map_path=None,
          cafe24_path=None, cat_path=None, cafe24_basis="order",
          cafe24_net=None):
    period = gw.date_range_from_naver_name(naver_path)

    naver_total, naver_products = gw.parse_naver(naver_path)
    coupang_total, coupang_products = gw.parse_coupang(coupang_path)
    ohouse_total, ohouse_products = gw.parse_ohouse(ohouse_path, password=ohouse_pw)
    # 자사몰: 주간 엔진의 parse_cafe24 → 공용 `cafe24.py`(엔진 폴더)로 위임된다.
    # 제품명 I열 / 옵션 J열 / 금액 F열(총 주문금액), G열 0원은 취소 처리. 주간과 동일 규칙.
    cafe24_total, cafe24_products = (
        gw.cafe24.parse(cafe24_path, gw.group_and_option, basis=cafe24_basis)
        if cafe24_path else (0.0, []))

    if base_path and os.path.exists(base_path):
        wb = openpyxl.load_workbook(base_path)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "매출데이터"

    def sheet(name):
        """항상 새로 만든다(매달 덮어씀). 시트 위치는 유지."""
        if name in wb.sheetnames:
            idx = wb.sheetnames.index(name)
            wb.remove(wb[name])
            return wb.create_sheet(name, idx)
        return wb.create_sheet(name)

    if "매출데이터" not in wb.sheetnames:
        wb.create_sheet("매출데이터", 0)
    # 자사몰 매출데이터 값: 카페24 관리자 '월별 매출'의 **순매출**이 정답이다.
    # 주문내역 CSV에는 환불 정보가 없어(전액취소만 G=0으로 보임) 순매출을 만들 수 없다.
    # --cafe24-net 으로 받은 값이 있으면 그것을 쓰고, 없으면 CSV 집계값을 쓴다.
    cafe24_for_summary = cafe24_net if cafe24_net is not None else (
        cafe24_total if cafe24_path else None)
    totals = {"네이버": naver_total, "쿠팡": coupang_total, "오늘의집": ohouse_total,
              "자사몰": cafe24_for_summary}   # 파일도 값도 없으면 빈칸
    months, chans = update_summary(wb["매출데이터"], month_label, totals, brand)

    # 통합순위 채널 = 매출데이터에 표기된 채널 중 이번 달 순위 데이터가 있는 것
    plists = {"네이버": naver_products, "쿠팡": coupang_products,
              "오늘의집": ohouse_products, "자사몰": cafe24_products}
    uchans = [c for c in chans if plists[c] and c not in UNIFIED_EXCLUDE]
    attach_quantities(gw, naver_products, coupang_products, ohouse_products,
                      naver_path, ohouse_path, ohouse_pw)
    unified, csv_rows = mp.build_unified(
        [(c, plists[c]) for c in uchans],
        gw._is_color, gw._is_size, map_path=map_path, cat_path=cat_path)
    if map_path:
        mp.save_map(map_path, csv_rows)

    write_unified_sheet(sheet("통합순위"), period, unified, uchans)
    gw.write_naver_sheet(sheet("네이버"), period, naver_products)
    gw.write_coupang_sheet(sheet("쿠팡"), period, coupang_products)
    gw.write_ohouse_sheet(sheet("오늘의집"), period, ohouse_products)
    if cafe24_products:
        gw.write_cafe24_sheet(sheet("자사몰"), period, cafe24_products)

    desired = ["매출데이터", "통합순위", "네이버", "쿠팡", "오늘의집", "자사몰"]
    front = [wb[n] for n in desired if n in wb.sheetnames]
    wb._sheets = front + [w for w in wb.worksheets if w.title not in desired]
    wb.save(out_path)

    total = sum(totals[c] or 0 for c in chans)
    merged = sum(1 for r in unified if r["채널수"] >= 2)
    print(f"[완료] {out_path}")
    print(f"  기간: {period}   월: {month_label}   회사: {brand}")
    print("  " + " | ".join(f"{c} {int(totals[c]):,}" if totals[c] is not None else f"{c} -" for c in chans))
    if cafe24_net is not None and cafe24_path:
        print(f"  (자사몰: 매출데이터 {int(cafe24_net):,} = 카페24 순매출 / "
              f"순위표 합 {int(cafe24_total):,} = 주문 기준 — 환불 반영 차이)")
    print(f"  합계  {int(total):,}")
    print(f"  누적 월: {', '.join(months)}")
    print("  순위 제품수 — " + " / ".join(f"{c} {len(plists[c])}" for c in uchans) +
          f" → 통합 {len(unified)}종 (2개 채널 이상 {merged}종)")
    if map_path:
        print(f"  매칭표: {map_path}  (통합제품명 열을 고치면 다음 실행에 반영됨)")
    if cat_path:
        print(f"  카테고리표: {cat_path}  (카테고리 열을 고치면 다음 실행에 반영됨)")
    return out_path


def main():
    ap = argparse.ArgumentParser(description="월간 판매 데이터 자동 생성기")
    ap.add_argument("--input", help="월 단위 엑셀 3개(+지난달 결과)가 든 폴더")
    ap.add_argument("--naver")
    ap.add_argument("--coupang")
    ap.add_argument("--ohouse")
    ap.add_argument("--cafe24", help="자사몰(카페24) 주문내역 CSV")
    ap.add_argument("--cafe24-basis", choices=["order", "paid"], default="order",
                    help="자사몰 순위표 금액 기준: order=F열 총 주문금액(기본) / paid=G열 총 결제금액")
    ap.add_argument("--cafe24-net", type=int,
                    help="자사몰 월매출(매출데이터용). 카페24 관리자 "
                         "'통계>매출분석>월별 매출'의 **순매출** 값을 넣는다. "
                         "CSV에는 환불이 안 담겨 이 값을 만들 수 없다.")
    ap.add_argument("--no-cafe24", action="store_true", help="자사몰 데이터를 쓰지 않음")
    ap.add_argument("--ohouse-pw", default=None)
    ap.add_argument("--month", help="예: 8월 (미지정 시 네이버 파일명에서 산출)")
    ap.add_argument("--base", help="지난달 월간판매데이터.xlsx (매출데이터 누적)")
    ap.add_argument("--brand", default="플레르가구", help="플레르가구 / 가즈다가구")
    ap.add_argument("--out", required=True)
    ap.add_argument("--engine", help="generate_weekly_sales.py 경로 직접 지정")
    ap.add_argument("--map", help="통합 매칭표 CSV 경로 (기본: 통합_제품맵_{회사}.csv)")
    ap.add_argument("--cat-map", help="카테고리 수정표 CSV 경로 (기본: 통합_카테고리_{회사}.csv)")
    ap.add_argument("--no-map", action="store_true", help="매칭표 CSV를 쓰지도 읽지도 않음")
    a = ap.parse_args()

    gw = load_engine(a.engine)
    pw = a.ohouse_pw or gw.DEFAULT_OHOUSE_PW

    naver, coupang, ohouse, base = a.naver, a.coupang, a.ohouse, a.base
    cafe24 = a.cafe24
    if a.input:
        f = gw.classify_files(a.input, ohouse_pw=pw)
        naver = naver or f["naver"]
        coupang = coupang or f["coupang"]
        ohouse = ohouse or f["ohouse"]
        base = base or f["base"]
        cafe24 = cafe24 or f.get("cafe24")
    if a.no_cafe24:
        cafe24 = None

    missing = [n for n, v in [("네이버", naver), ("쿠팡", coupang), ("오늘의집", ohouse)] if not v]
    if missing:
        raise SystemExit(
            "다음 플랫폼 파일을 인식하지 못했습니다: " + ", ".join(missing) +
            "\n  - 네이버: 'sales_YYYYMMDD-YYYYMMDD_*.xlsx' (상품별 통계, 시트 SALES)"
            "\n  - 쿠팡: 'SELLER_INSIGHTS_VENDOR_ITEM_METRICS_*.xlsx'"
            "\n  - 오늘의집: '주문배송 내역*.xlsx' (암호 파일)")

    month = a.month or month_from_naver_name(naver)
    if not month:
        raise SystemExit("월을 알 수 없습니다. --month \"8월\" 로 지정하세요.")

    map_path = cat_path = None
    if not a.no_map:
        outdir = os.path.dirname(os.path.abspath(a.out)) or HERE
        map_path = a.map or os.path.join(outdir, f"통합_제품맵_{a.brand}.csv")
        cat_path = a.cat_map or os.path.join(outdir, f"통합_카테고리_{a.brand}.csv")

    build(gw, naver, coupang, ohouse, a.out, month,
          ohouse_pw=pw, base_path=base, brand=a.brand, map_path=map_path,
          cafe24_path=cafe24, cat_path=cat_path, cafe24_basis=a.cafe24_basis,
          cafe24_net=a.cafe24_net)


if __name__ == "__main__":
    sys.exit(main())
